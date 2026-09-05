import csv
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from migration_report_tool.infrastructure.database.sqlite_store import ProjectStore
from migration_report_tool.infrastructure.database.migrations import TARGET_SCHEMA_VERSION
from migration_report_tool.services.derived_table_service import (
    build_derived_table,
    evaluate_expression,
    export_derived_csv,
    export_derived_xlsx,
    source_field_catalog,
)


class DerivedTableBuilderTests(unittest.TestCase):
    def test_safe_expression_functions(self):
        context = {
            "adms_db": {"rmu": "17233", "smart": "SMART"},
            "zenon_sld": {"feeder": "JED-NTH-ABH-15"},
        }
        self.assertEqual(evaluate_expression("adms_db.rmu", context), "17233")
        self.assertEqual(evaluate_expression("COALESCE('', adms_db.smart)", context), "SMART")
        self.assertEqual(
            evaluate_expression('IF(NORMALIZE(zenon_sld.feeder) == "JEDNTHABH15", "Closed", "Needs Action")', context),
            "Closed",
        )
        with self.assertRaises(ValueError):
            evaluate_expression("__import__('os').system('echo bad')", context)

    def test_custom_fields_join_calculate_and_export(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            store = ProjectStore(root / "project")
            try:
                adms = root / "ADMS-DB.csv"
                adms.write_text(
                    "RMU_NAME,ADMS_GSS-FID,SMART,VOLT\n"
                    "1,JED-NTH-ABH-01,SMART,13.8\n"
                    "2,JED-NTH-ABH-02,NORMAL,13.8\n",
                    encoding="utf-8",
                )
                zenon = root / "ZENON-SLD.csv"
                zenon.write_text(
                    "RMU,Feeder,CabinetType\n"
                    "1,JED-NTH-ABH-01,2L1T\n"
                    "2,JED-NTH-WRONG-99,2L1T\n",
                    encoding="utf-8",
                )
                store.replace_custom_source_fields(
                    "adms_db",
                    [{"field_key": "voltage", "display_name": "Voltage", "actual_column": "VOLT"}],
                    "tester",
                )
                catalog = source_field_catalog(store, "adms_db")
                self.assertTrue(any(item["key"] == "voltage" and item["custom"] for item in catalog))

                config = {
                    "table_name": "RMU_FINAL",
                    "description": "test",
                    "base_source_type": "adms_db",
                    "joins": [
                        {
                            "source_type": "zenon_sld",
                            "join_type": "LEFT",
                            "left_expression": "adms_db.rmu",
                            "right_field": "rmu",
                        }
                    ],
                    "fields": [
                        {"name": "RMU", "key": "rmu", "expression": "adms_db.rmu"},
                        {"name": "Voltage", "key": "voltage", "expression": "adms_db.voltage"},
                        {
                            "name": "Review",
                            "key": "review",
                            "expression": 'IF(NORMALIZE(zenon_sld.feeder) == NORMALIZE(adms_db.gss_fid), "Closed", "Needs Action")',
                        },
                    ],
                }
                store.save_derived_table_config(config, "tester")
                saved = store.derived_table_config("RMU_FINAL")
                self.assertIsNotNone(saved)
                self.assertEqual(saved["config"]["fields"][1]["expression"], "adms_db.voltage")

                result = build_derived_table(
                    saved["config"],
                    store,
                    {"adms_db": adms, "zenon_sld": zenon},
                )
                self.assertEqual(result.columns, ("RMU", "Voltage", "Review"))
                self.assertEqual(len(result.rows), 2)
                self.assertEqual(result.rows[0]["Voltage"], "13.8")
                self.assertEqual(result.rows[0]["Review"], "Closed")
                self.assertEqual(result.rows[1]["Review"], "Needs Action")

                csv_path = export_derived_csv(result, root / "out.csv")
                with csv_path.open("r", encoding="utf-8-sig", newline="") as handle:
                    rows = list(csv.DictReader(handle))
                self.assertEqual(rows[1]["Review"], "Needs Action")

                xlsx_path = export_derived_xlsx(result, root / "out.xlsx")
                wb = load_workbook(xlsx_path, read_only=True, data_only=True)
                try:
                    ws = wb["DATA"]
                    self.assertEqual(ws.cell(1, 1).value, "RMU")
                    self.assertEqual(ws.cell(2, 3).value, "Closed")
                finally:
                    wb.close()
            finally:
                store.close()

    def test_schema_v5_tables_exist(self):
        with tempfile.TemporaryDirectory() as td:
            store = ProjectStore(Path(td) / "site")
            try:
                self.assertEqual(store.project_schema_version(), TARGET_SCHEMA_VERSION)
                names = {
                    row[0] for row in store.db.execute(
                        "SELECT name FROM sqlite_master WHERE type='table'"
                    )
                }
                self.assertIn("custom_source_fields", names)
                self.assertIn("derived_table_configs", names)
            finally:
                store.close()


if __name__ == "__main__":
    unittest.main()
