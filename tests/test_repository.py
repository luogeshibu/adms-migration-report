from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from migration_report_tool.repository import resolve_site_sources, scan_repository


class SiteRepositoryTests(unittest.TestCase):
    def test_unversioned_canonical_names_are_auto_detection_fallbacks(self):
        with tempfile.TemporaryDirectory() as td:
            site = Path(td) / "ADF"
            site.mkdir()
            for name in [
                "SE.xlsx", "ZENON-DB.csv", "ZENON-SLD.csv",
                "ADMS-DB.csv", "ADMS-SLD.csv", "ZENON-ADMS-IOA.csv",
            ]:
                (site / name).write_bytes(b"x")
            sources = resolve_site_sources(site)
            self.assertEqual(sources["zenon_sld"].name, "ZENON-SLD.csv")
            self.assertEqual(sources["adms_db"].name, "ADMS-DB.csv")
            info = scan_repository(Path(td))[0]
            self.assertTrue(info.ready)
            self.assertEqual(info.detections["zenon_sld"].method, "Auto Detection Rule")

    def test_current_legacy_unversioned_names_use_rules_when_no_v_release_exists(self):
        with tempfile.TemporaryDirectory() as td:
            site = Path(td) / "ADF"
            site.mkdir()
            for name in ["ADF-SE.xlsx", "ZENON-DB.csv", "ZENON-SLD.csv", "ADMS-DB.csv", "ADMS-SLD.csv", "ZENON-ADMS-IOA.csv"]:
                (site / name).write_bytes(b"x")
            sources = resolve_site_sources(site)
            self.assertIn("se_list", sources)
            self.assertTrue(scan_repository(Path(td))[0].ready)

    def test_zenon_sld_csv_is_required_and_xml_files_are_ignored(self):
        with tempfile.TemporaryDirectory() as td:
            site = Path(td) / "ADF"
            site.mkdir()
            for name in ["SE.xlsx", "ZENON-DB.csv", "ADMS-DB.csv", "ADMS-SLD.csv", "ZENON-ADMS-IOA.csv"]:
                (site / name).write_bytes(b"x")
            (site / "ADF.XML").write_bytes(b"<Subject/>")
            info = scan_repository(Path(td), deep=False)[0]
            self.assertFalse(info.ready)
            self.assertIn("zenon_sld", info.missing_required)
            self.assertNotIn("zenon_xml", info.sources)
            self.assertNotIn("ADF.XML", [p.name for p in info.unmapped_files])

    def test_ioa_is_required_for_formal_delivery_readiness(self):
        with tempfile.TemporaryDirectory() as td:
            site = Path(td) / "ADF"
            site.mkdir()
            for name in ["SE.xlsx", "ZENON-DB.csv", "ZENON-SLD.csv", "ADMS-DB.csv", "ADMS-SLD.csv"]:
                (site / name).write_bytes(b"x")
            info = scan_repository(Path(td))[0]
            self.assertFalse(info.ready)
            self.assertIn("ioa", info.missing_required)

    def test_site_report_workbook_is_not_a_repository_source(self):
        with tempfile.TemporaryDirectory() as td:
            site = Path(td) / "ADF"
            site.mkdir()
            for name in ["SE.xlsx", "ZENON-DB.csv", "ADMS-DB.csv", "ADMS-SLD.csv", "ADF-REPORT.xlsx"]:
                (site / name).write_bytes(b"x")
            sources = resolve_site_sources(site)
            self.assertNotIn("report_template", sources)
            self.assertNotIn("standard_reference", sources)


if __name__ == "__main__":
    unittest.main()


def test_comparison_column_layout_migration_adds_ip_and_retires_link():
    from migration_report_tool.schema import (
        COMPARISON_COLUMN_SCHEMA_VERSION,
        migrate_comparison_visible_columns,
    )

    old_saved = [
        "no", "rmu", "analysis_name", "analysis_feeder",
        "analysis_smart", "analysis_type", "remarks", "comments",
    ]
    migrated = migrate_comparison_visible_columns(old_saved, 1)
    assert "analysis_ip" in migrated
    assert "analysis_link" not in migrated

    # Once the schema version is current, an explicit user hide is preserved.
    current_saved = [key for key in migrated if key != "analysis_ip"]
    preserved = migrate_comparison_visible_columns(
        current_saved, COMPARISON_COLUMN_SCHEMA_VERSION
    )
    assert "analysis_ip" not in preserved
    assert "analysis_link" not in preserved

class SourceDiscoveryV0838Tests(unittest.TestCase):
    def test_se_equipment_can_be_detected_by_content_with_arbitrary_filename(self):
        from openpyxl import Workbook
        import migration_report_tool.infrastructure.filesystem.site_repository as repo
        from unittest.mock import patch
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            site = root / "ADEL"
            site.mkdir()
            config_path = root / "site_repository.json"
            wb = Workbook()
            ws = wb.active
            ws.append(["SS", "FEEDER", "EQUIPMENT", "EQUIP. TYPE", "OH / UG"])
            ws.append(["ADEL", "ADEL-01", "1001", "SMART", "UG"])
            wb.save(site / "ADEL-V1.xlsx")
            wb.close()
            wb = Workbook()
            wb.active.append(["Report", "Value"])
            wb.save(site / "unrelated.xlsx")
            wb.close()
            with patch.object(repo, "_config_path", return_value=config_path):
                info = repo.scan_repository(root)[0]
            self.assertEqual(info.sources["se_list"].name, "ADEL-V1.xlsx")
            self.assertEqual(info.detections["se_list"].method, "Content")
            self.assertGreaterEqual(info.detections["se_list"].confidence, 75)
            self.assertIn("unrelated.xlsx", [p.name for p in info.unmapped_files])

    def test_xml_files_are_outside_source_detection(self):
        import migration_report_tool.infrastructure.filesystem.site_repository as repo
        from unittest.mock import patch
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            site = root / "ABH"
            site.mkdir()
            config_path = root / "site_repository.json"
            (site / "custom-V1.xml").write_text("<Subject/>", encoding="utf-8")
            with patch.object(repo, "_config_path", return_value=config_path):
                info = repo.scan_repository(root, deep=False)[0]
            self.assertNotIn("zenon_xml", info.sources)
            self.assertNotIn("custom-V1.xml", [p.name for p in info.unmapped_files])

class IncrementalSourceImportV0839Tests(unittest.TestCase):
    def test_rank_source_candidates_identifies_arbitrary_se_workbook(self):
        from openpyxl import Workbook
        from migration_report_tool.repository import rank_source_candidates
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "现场设备清单.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.append(["SS", "FEEDER", "EQUIPMENT", "EQUIP. TYPE", "OH / UG"])
            ws.append(["ABH", "ABH-01", "1001", "SMART", "UG"])
            wb.save(path)
            wb.close()
            ranked = rank_source_candidates(path)
            self.assertTrue(ranked)
            self.assertEqual(ranked[0][0], "se_list")
            self.assertGreaterEqual(ranked[0][1], 75)

    def test_workspace_manual_override_survives_repository_sync(self):
        from migration_report_tool.repository import SiteInfo, source_status, sync_site_to_project
        from migration_report_tool.storage import ProjectStore
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            site_dir = root / "repo" / "ABH"
            site_dir.mkdir(parents=True)
            repository_csv = site_dir / "ZENON-DB.csv"
            repository_csv.write_text("repo", encoding="utf-8")
            manual_csv = root / "manual.csv"
            manual_csv.write_text("manual", encoding="utf-8")
            store = ProjectStore(root / "workspace")
            try:
                store.config["repository_site"] = "ABH"
                store.config["repository_path"] = str(site_dir)
                store.save_config()
                stored_manual = store.set_source("zenon_db", manual_csv)
                store.mark_manual_source_override("zenon_db", manual_csv)
                site = SiteInfo("ABH", site_dir, {"zenon_db": repository_csv})
                status = source_status(site, store)["zenon_db"]
                self.assertEqual(status["status"], "MANUAL")
                self.assertEqual(Path(status["path"]), stored_manual)
                sync_site_to_project(store, site)
                self.assertEqual(store.source_path("zenon_db"), stored_manual)
                self.assertEqual(store.source_path("zenon_db").read_text(encoding="utf-8"), "manual")
            finally:
                store.db.close()


class SourceVersionSelectionV0857Tests(unittest.TestCase):
    def _write_zenon_sld(self, path: Path, feeder: str):
        path.write_text(
            "RMU,Feeder,CabinetType\n10689,%s,3L1T\n" % feeder,
            encoding="utf-8",
        )

    def test_auto_selects_highest_explicit_v_and_ignores_unversioned_base(self):
        from migration_report_tool.repository import resolve_site_sources, source_file_version
        with tempfile.TemporaryDirectory() as td:
            site = Path(td) / "1-ABH"
            site.mkdir()
            self._write_zenon_sld(site / "ZENON-SLD.csv", "BASE")
            self._write_zenon_sld(site / "ZENON-SLD-V1.csv", "V1")
            self._write_zenon_sld(site / "ZENON-SLD-V2.1.0.csv", "V210")
            sources = resolve_site_sources(site)
            self.assertEqual(sources["zenon_sld"].name, "ZENON-SLD-V2.1.0.csv")
            self.assertEqual(source_file_version(sources["zenon_sld"]), (2, 1, 0))

    def test_unversioned_file_is_detection_rule_fallback_when_no_published_v_exists(self):
        from migration_report_tool.repository import resolve_site_sources
        with tempfile.TemporaryDirectory() as td:
            site = Path(td) / "1-ABH"
            site.mkdir()
            self._write_zenon_sld(site / "ZENON-SLD.csv", "BASE")
            sources = resolve_site_sources(site)
            self.assertEqual(sources["zenon_sld"].name, "ZENON-SLD.csv")

    def test_project_pinned_source_version_overrides_auto_latest_and_survives_reopen(self):
        from migration_report_tool.repository import scan_repository, selected_repository_source_path, sync_site_to_project
        from migration_report_tool.storage import ProjectStore
        with tempfile.TemporaryDirectory() as td:
            root = Path(td) / "repo"
            site_dir = root / "1-ABH"
            site_dir.mkdir(parents=True)
            self._write_zenon_sld(site_dir / "ZENON-SLD-V1.csv", "V1")
            self._write_zenon_sld(site_dir / "ZENON-SLD-V2.csv", "V2")
            site = scan_repository(root)[0]
            self.assertEqual(site.sources["zenon_sld"].name, "ZENON-SLD-V2.csv")

            project = Path(td) / "project"
            store = ProjectStore(project)
            try:
                store.config["repository_site"] = site.name
                store.config["repository_path"] = str(site.path)
                store.save_config()
                store.set_source_file_selection("zenon_sld", "ZENON-SLD-V1.csv", "reviewer")
                self.assertEqual(selected_repository_source_path(site, store, "zenon_sld").name, "ZENON-SLD-V1.csv")
                sync_site_to_project(store, site)
                self.assertEqual(store.source_path("zenon_sld").read_text(encoding="utf-8").splitlines()[1].split(",")[1], "V1")
            finally:
                store.close()

            reopened = ProjectStore(project)
            try:
                self.assertEqual(reopened.source_file_selection("zenon_sld")["file_name"], "ZENON-SLD-V1.csv")
            finally:
                reopened.close()
