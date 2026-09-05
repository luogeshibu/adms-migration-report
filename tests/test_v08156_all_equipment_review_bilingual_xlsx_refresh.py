from __future__ import annotations

import os
import time
from pathlib import Path

from openpyxl import Workbook

from migration_report_tool.adapters import SourceAdapter
from migration_report_tool.infrastructure.database.sqlite_store import ProjectStore
from migration_report_tool.infrastructure.filesystem.site_repository import (
    SiteInfo, discover_site_sources, sync_site_to_project,
)
from migration_report_tool.services.rmu_review_service import build_equipment_source_view
from migration_report_tool.version import __version__


class MultiAdapter(SourceAdapter):
    def __init__(self, tables):
        self.tables = {k: list(v) for k, v in tables.items()}

    def load_rows(self, source_type: str):
        return [dict(row) for row in self.tables.get(source_type, [])]


def _zenon_sld_book(path: Path, *, name: str = "L100", device_type: str = "LBS"):
    wb = Workbook()
    ws = wb.active
    ws.title = "DEVICES"
    ws.append(["DeviceName", "DeviceScope", "SubType", "SMART", "DeviceType", "Picture"])
    ws.append([name, "JED-STH-ADEL-20", "OLBS", "SMART", device_type, "ADEL-110"])
    wb.save(path)


def test_release_version():
    assert __version__ == "0.8.156"


def test_non_rmu_uses_same_five_source_analysis_and_namespaced_review_key(tmp_path: Path):
    tables = {
        "zenon_sld": [{"device_type": "LBS", "rmu": "1196", "feeder": "JED-STH-ADEL-20", "cabinet_type": "OLBS", "smart": "SMART"}],
        "se_list": [{"rmu": "1196", "feeder": "ADEL-20", "rmu_type": "OLBS", "smart": "SMART"}],
        "zenon_db": [{"rmu": "1196", "feeder": "ADEL-20", "rmu_type": "OLBS", "smart": "SMART", "ip": "10.0.0.1"}],
        "adms_db": [{"rmu": "1196", "gss_fid": "ADEL-20", "rmu_type": "OLBS", "smart": "NORMAL", "channel_ip": "10.0.0.2"}],
        "adms_sld": [{"rmu": "1196", "feeder": "ADEL-20", "rmu_type": "OLBS", "smart": "SMART"}],
    }
    store = ProjectStore(tmp_path / "project")
    try:
        rows, _ = build_equipment_source_view(store, "LBS", MultiAdapter(tables))
        assert len(rows) == 1
        row = rows[0]
        assert row["review_key"] == "EQ::LBS::1196"
        assert row["equipment_source_count"] == "5/5"
        assert row["analysis_name"] == "TRUE"
        assert row["analysis_feeder"] == "TRUE"
        assert row["analysis_type"] == "TRUE"
        assert row["analysis_smart"] == "FALSE"
        assert row["analysis_ip"] == "FALSE"
    finally:
        store.close()


def test_existing_rmu_review_and_new_equipment_review_are_independent_and_persistent(tmp_path: Path):
    folder = tmp_path / "project"
    store = ProjectStore(folder)
    try:
        store.append_rmu_manual_review_comment("5977", "existing RMU comment", "tester")
        store.update_rmu_review_status("5977", "NEEDS ACTION", "tester", "legacy RMU review")
        store.append_rmu_manual_review_comment("EQ::LBS::1196", "new LBS comment", "tester")
        store.update_rmu_review_status("EQ::LBS::1196", "NEEDS ACTION", "tester", "LBS review")
    finally:
        store.close()

    reopened = ProjectStore(folder)
    try:
        review_map = reopened.rmu_review_map()
        assert review_map["5977"]["manual_comment"] == "existing RMU comment"
        assert review_map["5977"]["review_status"] == "NEEDS ACTION"
        assert review_map["EQ::LBS::1196"]["manual_comment"] == "new LBS comment"
        assert review_map["EQ::LBS::1196"]["review_status"] == "NEEDS ACTION"
        assert reopened.schema_version == 12
    finally:
        reopened.close()


def test_auto_refresh_prefers_newer_xlsx_over_older_csv_same_source_family(tmp_path: Path):
    site = tmp_path / "1-ADF"
    site.mkdir()
    csv_path = site / "ZENON-SLD.csv"
    csv_path.write_text("DeviceName,DeviceScope,SubType,SMART,DeviceType,Picture\n1001,JED-CTL-ADF-16,2L1T,SMART,RMU,ADF110\n", encoding="utf-8")
    # Ensure the workbook has a newer physical mtime.
    time.sleep(0.02)
    xlsx_path = site / "ZENON-SLD.xlsx"
    _zenon_sld_book(xlsx_path, name="1002", device_type="RMU")
    now = time.time_ns()
    os.utime(csv_path, ns=(now - 5_000_000_000, now - 5_000_000_000))
    os.utime(xlsx_path, ns=(now, now))

    sources, _detections, _unmapped = discover_site_sources(site, deep=False)
    assert sources["zenon_sld"].name == "ZENON-SLD.xlsx"


def test_snapshot_preserves_workbook_extension(tmp_path: Path):
    site_dir = tmp_path / "repo" / "1-ADF"
    site_dir.mkdir(parents=True)
    workbook = site_dir / "ZENON-SLD.xlsx"
    _zenon_sld_book(workbook)
    site = SiteInfo("1-ADF", site_dir, sources={"zenon_sld": workbook})
    store = ProjectStore(tmp_path / "project")
    try:
        result = sync_site_to_project(store, site, force_all=True)
        manifest_xlsx = result.snapshot_dir / "ZENON-SLD.xlsx"
        assert manifest_xlsx.exists()
        assert not (result.snapshot_dir / "ZENON-SLD.csv").exists()
        assert store.source_path("zenon_sld").suffix.lower() == ".xlsx"
    finally:
        store.close()


def test_critical_chinese_ui_strings_are_covered_without_translating_physical_headers():
    text = Path("src/migration_report_tool/ui/i18n.py").read_text(encoding="utf-8")
    for source, chinese in {
        "Project Overview": "项目总览",
        "Site Data Sources": "站点数据源",
        "Equipment Data Review": "设备数据审核",
        "Search site...": "搜索站点...",
        "Current Migration Site": "当前迁移站点",
        "Active Data Sources": "当前数据源",
        "Review Equipment Issues": "审核设备问题",
        "MAPPING REQUIRED": "需要字段映射",
    }.items():
        assert repr(source)[1:-1] in text or source in text
        assert chinese in text
    assert '"DeviceName":' not in text
