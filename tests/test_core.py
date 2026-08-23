import unittest
import tempfile
from pathlib import Path
try:
    from .selftest_core import run_selftest
except ImportError:  # unittest discover -s tests loads test_core as a top-level module
    from selftest_core import run_selftest
from migration_report_tool.parsers import parse_zenon_xml, feeder_matches_site
from migration_report_tool.storage import ProjectStore
from migration_report_tool.comparison import build_comparison
from migration_report_tool.importing import import_source
from openpyxl import Workbook

class CoreRegressionTest(unittest.TestCase):
    def test_end_to_end_core(self):
        run_selftest()

    def test_zenon_xml_picture_shortname_maps_to_screen_name(self):
        xml = """<?xml version='1.0' encoding='utf-8'?>
<Root>
  <Picture ShortName='ADF110 (JEDDAH)'>
    <Elements_Test>
      <LinkName>01_VERT_SRMU_2L1T_01</LinkName>
      <SubstituteDestination>ADF-15-5</SubstituteDestination>
    </Elements_Test>
  </Picture>
</Root>"""
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / 'sample.xml'
            path.write_text(xml, encoding='utf-8')
            rows = parse_zenon_xml(path)
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]['screen_name'], 'ADF110 (JEDDAH)')
        self.assertEqual(rows[0]['feeder'], 'ADF-15')
        self.assertEqual(rows[0]['rmu'], '5')

    def test_official_se_feeder_header_maps_to_report(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            source = root / "ADF.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.append(["SS", "FEEDER", "EQUIPMENT", "EQUIP. TYPE", "OH / UG"])
            ws.append(["ADF", "ADF-15", "7324B", "SMART NOP HT", "UG"])
            ws.append(["ADF", "ADF-16", 26859, "SMART NOP HT", "UG"])
            wb.save(source)
            wb.close()
            store = ProjectStore(root / "project")
            try:
                store.set_source("se_list", source)
                rows, _ = build_comparison(store)
                by_rmu = {r["rmu"]: r for r in rows}
                self.assertEqual(by_rmu["7324B"]["se_feeder"], "ADF-15")
                self.assertEqual(by_rmu["26859"]["se_feeder"], "ADF-16")
                self.assertEqual(by_rmu["26859"]["se_station"], "ADF")
            finally:
                store.db.close()

    def test_reimported_se_file_becomes_latest_source(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            first = root / "ADF_first.xlsx"
            second = root / "ADF_second.xlsx"
            for path, feeder in ((first, "ADF-15"), (second, "ADF-99")):
                wb = Workbook()
                ws = wb.active
                ws.append(["SS", "FEEDER", "EQUIPMENT", "EQUIP. TYPE", "OH / UG"])
                ws.append(["ADF", feeder, "7324B", "SMART NOP HT", "UG"])
                wb.save(path)
                wb.close()
            store = ProjectStore(root / "project")
            try:
                r1 = import_source(store, "se_list", first)
                rows, _ = build_comparison(store)
                self.assertEqual(rows[0]["se_feeder"], "ADF-15")
                r2 = import_source(store, "se_list", second)
                rows, _ = build_comparison(store)
                self.assertEqual(rows[0]["se_feeder"], "ADF-99")
                self.assertNotEqual(r1.stored_path, r2.stored_path)
                self.assertEqual(store.source_path("se_list"), r2.stored_path)
            finally:
                store.db.close()

    def test_combined_abn_abn2_xml_isolated_by_exact_feeder_token(self):
        xml = """<?xml version='1.0' encoding='utf-8'?>
<Root>
  <Picture ShortName='ABN2-110 (JEDDAH)'>
    <Elements_0><LinkName>01_VERT_SRMU_2L1T_01</LinkName><SubstituteDestination>JED-NTH-ABN2-16-24668</SubstituteDestination></Elements_0>
    <Elements_1><LinkName>01_VERT_SRMU_2L1T_01</LinkName><SubstituteDestination>JED-NTH-ABN-22-23953</SubstituteDestination></Elements_1>
  </Picture>
</Root>"""
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "ABN-ABN2.XML"
            path.write_text(xml, encoding="utf-8")
            abn2 = parse_zenon_xml(path, site_name="ABN2")
            abn = parse_zenon_xml(path, site_name="ABN")
        self.assertEqual([(r["feeder"], r["rmu"]) for r in abn2], [("JED-NTH-ABN2-16", "24668")])
        self.assertEqual([(r["feeder"], r["rmu"]) for r in abn], [("JED-NTH-ABN-22", "23953")])
        self.assertFalse(feeder_matches_site("JED-NTH-ABN2-16", "ABN"))
        self.assertTrue(feeder_matches_site("JED-NTH-ABN2-16", "ABN2"))

    def test_se_feeder_suffix_is_an_additional_positive_filter(self):
        self.assertTrue(feeder_matches_site("JED-NTH-ABN2-16", "CUSTOM_SITE", ["ABN2-16"]))
        self.assertFalse(feeder_matches_site("JED-NTH-ABN-16", "CUSTOM_SITE", ["ABN2-16"]))


if __name__ == '__main__': unittest.main()

class AnalysisConsistencyTests(unittest.TestCase):
    def test_shared_consistency_rule(self):
        from migration_report_tool.analysis import compare_consistency, normalize_name
        self.assertIsNone(compare_consistency({"A": "", "B": None}, normalize_name).value)
        self.assertTrue(compare_consistency({"A": "26859", "B": ""}, normalize_name).value)
        self.assertTrue(compare_consistency({"A": "26859", "B": "26859"}, normalize_name).value)
        self.assertFalse(compare_consistency({"A": "26859", "B": "26860"}, normalize_name).value)

    def test_feeder_normalization_keeps_station_and_ignores_region_prefix(self):
        from migration_report_tool.analysis import normalize_feeder_for_compare
        values = {
            normalize_feeder_for_compare("ABN2-03", "ABN2"),
            normalize_feeder_for_compare("JED-NTH-ABN2-3", "ABN2"),
            normalize_feeder_for_compare("JED NTH-ABN2-03", "ABN2"),
        }
        self.assertEqual(values, {"ABN2-3"})

    def test_abh_feeder_normalization_decodes_adms_ah3_encoding(self):
        from migration_report_tool.analysis import normalize_feeder_for_compare
        equivalents_03 = {
            normalize_feeder_for_compare("ABH-03", "ABH"),
            normalize_feeder_for_compare("ABH-3", "ABH"),
            normalize_feeder_for_compare("JED-NTH-ABH-3", "ABH"),
            normalize_feeder_for_compare("JED-NTH-ABH-AH303", "ABH"),
        }
        self.assertEqual(equivalents_03, {"ABH-3"})
        self.assertEqual(normalize_feeder_for_compare("JED-NTH-ABH-AH308", "ABH"), "ABH-8")
        # Repository/site labels are not guaranteed to equal the feeder site token.
        self.assertEqual(normalize_feeder_for_compare("JED-NTH-ABH-AH303", "JEDDAH_ABH"), "ABH-3")
        self.assertEqual(normalize_feeder_for_compare("JED-NTH-ABH-3", "ABH110"), "ABH-3")
        self.assertEqual(normalize_feeder_for_compare("JED-NTH-ABH-AH303", None), "ABH-3")

    def test_abn_adms_ah3_feeder_encoding_uses_logical_suffix_number(self):
        from migration_report_tool.analysis import normalize_feeder_for_compare
        self.assertEqual(normalize_feeder_for_compare("ABN-12", "1-ABN"), "ABN-12")
        self.assertEqual(normalize_feeder_for_compare("JED-NTH-ABN-12", "1-ABN"), "ABN-12")
        self.assertEqual(normalize_feeder_for_compare("ABN-AH312", "1-ABN"), "ABN-12")
        self.assertEqual(normalize_feeder_for_compare("ABN-1", "1-ABN"), "ABN-1")
        self.assertEqual(normalize_feeder_for_compare("ABN-01", "1-ABN"), "ABN-1")
        self.assertEqual(normalize_feeder_for_compare("ABN-AH301", "1-ABN"), "ABN-1")

    def test_abn_feeder_consistency_matches_user_examples(self):
        from migration_report_tool.analysis import compare_consistency, normalize_feeder_for_compare
        normalizer = lambda value: normalize_feeder_for_compare(value, "1-ABN")
        result_12 = compare_consistency({
            "SE": "ABN-12",
            "ZENON DB": "ABN-12",
            "ZENON SLD XML": "JED-NTH-ABN-12",
            "ADMS DB": "ABN-AH312",
        }, normalizer)
        self.assertTrue(result_12.value)
        result_1 = compare_consistency({
            "SE": "ABN-1",
            "ZENON SLD XML": "JED-NTH-ABN-1",
            "ADMS DB": "ABN-AH301",
        }, normalizer)
        self.assertTrue(result_1.value)
        self.assertEqual(set(result_1.normalized_by_source.values()), {"ABN-1"})

    def test_unknown_abh_feeder_encoding_is_not_guessed(self):
        from migration_report_tool.analysis import normalize_feeder_for_compare
        self.assertEqual(normalize_feeder_for_compare("JED-NTH-ABH-DHN40", "ABH"), "ABH-DHN40")

    def test_abh_feeder_consistency_matches_user_examples(self):
        from migration_report_tool.analysis import compare_consistency, normalize_feeder_for_compare
        normalizer = lambda value: normalize_feeder_for_compare(value, "ABH")
        matching = compare_consistency({
            "SE": "ABH-8",
            "ZENON SLD XML": "JED-NTH-ABH-8",
            "ADMS DB": "JED-NTH-ABH-AH308",
        }, normalizer)
        self.assertTrue(matching.value)
        mismatch = compare_consistency({
            "SE": "ABH-03",
            "ZENON DB": "DHN-40",
            "ZENON SLD XML": "JED-NTH-ABH-3",
            "ADMS DB": "JED-NTH-ABH-AH303",
        }, normalizer)
        self.assertFalse(mismatch.value)
        self.assertEqual(mismatch.normalized_by_source["SE"], "ABH-3")
        self.assertEqual(mismatch.normalized_by_source["ZENON SLD XML"], "ABH-3")
        self.assertEqual(mismatch.normalized_by_source["ADMS DB"], "ABH-3")
        self.assertEqual(mismatch.normalized_by_source["ZENON DB"], "DHN-40")

    def test_same_number_different_station_is_feeder_mismatch(self):
        from migration_report_tool.analysis import compare_consistency, normalize_feeder_for_compare
        normalizer = lambda value: normalize_feeder_for_compare(value, "1-ABH")
        result = compare_consistency({
            "SE": "ABH-22",
            "ZENON DB": "JED-NTH-ABH-22",
            "ZENON SLD XML": "JED-NTH-ABH-22",
            "ADMS DB": "JED-NTH-ABN-22",
        }, normalizer)
        self.assertFalse(result.value)
        self.assertEqual(result.normalized_by_source["SE"], "ABH-22")
        self.assertEqual(result.normalized_by_source["ZENON DB"], "ABH-22")
        self.assertEqual(result.normalized_by_source["ZENON SLD XML"], "ABH-22")
        self.assertEqual(result.normalized_by_source["ADMS DB"], "ABN-22")

    def test_numeric_only_feeder_uses_selected_station_hint(self):
        from migration_report_tool.analysis import normalize_feeder_for_compare
        self.assertEqual(normalize_feeder_for_compare("22", "1-ABH"), "ABH-22")
        self.assertEqual(normalize_feeder_for_compare("22", "1-ABN"), "ABN-22")

    def test_zero_placeholder_handling(self):
        from migration_report_tool.analysis import normalize_feeder_for_compare, normalize_type, normalize_smart
        self.assertEqual(normalize_feeder_for_compare("0", "ABN2"), "")
        self.assertEqual(normalize_type("0"), "")
        self.assertEqual(normalize_smart("0"), "NORMAL")

    def test_smart_normalization(self):
        from migration_report_tool.analysis import normalize_smart
        self.assertEqual(normalize_smart("SMART NOP HT"), "SMART")
        self.assertEqual(normalize_smart("SMR"), "SMART")
        self.assertEqual(normalize_smart("YES"), "SMART")
        self.assertEqual(normalize_smart("NORMAL NOP"), "NORMAL")
        self.assertEqual(normalize_smart("NO"), "NORMAL")

    def test_end_to_end_analysis_uses_available_values_only(self):
        import csv
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            project = root / "project"
            se_path = root / "SE.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.append(["SS", "FEEDER", "EQUIPMENT", "EQUIP. TYPE", "OH / UG"])
            ws.append(["ABN2", "ABN2-03", "26859", "SMART NOP HT", "UG"])
            wb.save(se_path)
            wb.close()

            zdb_path = root / "ZENON-DB.csv"
            with zdb_path.open("w", encoding="utf-8-sig", newline="") as f:
                w = csv.DictWriter(f, fieldnames=["RMU", "FEEDER", "DEVICE", "SMART"])
                w.writeheader(); w.writerow({"RMU":"26859", "FEEDER":"ABN2-3", "DEVICE":"2L1T", "SMART":"SMART"})

            adb_path = root / "ADMS-DB.csv"
            with adb_path.open("w", encoding="utf-8-sig", newline="") as f:
                w = csv.DictWriter(f, fieldnames=["RMU_NAME", "ADMS_GSS-FID", "TYPE"])
                w.writeheader(); w.writerow({"RMU_NAME":"26859", "ADMS_GSS-FID":"JED NTH-ABN2-03", "TYPE":"2L1T"})

            asld_path = root / "ADMS-SLD.csv"
            with asld_path.open("w", encoding="utf-8-sig", newline="") as f:
                w = csv.DictWriter(f, fieldnames=["环网柜名称", "环网柜类型", "智能标识", "LINK"])
                w.writeheader(); w.writerow({"环网柜名称":"26859", "环网柜类型":"2L1T", "智能标识":"SMART", "LINK":"TRUE"})

            zsld_path = root / "ZENON-SLD.csv"
            with zsld_path.open("w", encoding="utf-8-sig", newline="") as f:
                w = csv.DictWriter(f, fieldnames=["RMU", "Feeder", "CabinetType"])
                w.writeheader(); w.writerow({"RMU":"26859", "Feeder":"JED-NTH-ABN2-03", "CabinetType":"2L1T"})

            store = ProjectStore(project)
            try:
                store.config["site_name"] = "ABN2"
                store.save_config()
                store.set_source("se_list", se_path)
                store.set_source("zenon_db", zdb_path)
                store.set_source("adms_db", adb_path)
                store.set_source("adms_sld", asld_path)
                store.set_source("zenon_sld", zsld_path)
                rows, _ = build_comparison(store)
                row = next(r for r in rows if r["rmu"] == "26859")
                self.assertEqual(row["analysis_name"], "TRUE")
                self.assertEqual(row["analysis_feeder"], "TRUE")
                self.assertEqual(row["analysis_smart"], "TRUE")
                self.assertEqual(row["analysis_type"], "TRUE")
            finally:
                store.db.close()

    def test_ip_and_link_analysis_keep_reason_details(self):
        import csv
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            zdb_path = root / "ZENON-DB.csv"
            with zdb_path.open("w", encoding="utf-8-sig", newline="") as f:
                w = csv.DictWriter(f, fieldnames=["RMU", "PRIMARY_IP", "PRIMARY_PORT"])
                w.writeheader(); w.writerow({"RMU":"R1", "PRIMARY_IP":"172.16.1.10", "PRIMARY_PORT":"2404"})
            adb_path = root / "ADMS-DB.csv"
            with adb_path.open("w", encoding="utf-8-sig", newline="") as f:
                w = csv.DictWriter(f, fieldnames=["RMU_NAME", "NET_DESCRIPTION1", "port"])
                w.writeheader(); w.writerow({"RMU_NAME":"R1", "NET_DESCRIPTION1":"172.16.1.11", "port":"2404"})
            asld_path = root / "ADMS-SLD.csv"
            with asld_path.open("w", encoding="utf-8-sig", newline="") as f:
                w = csv.DictWriter(f, fieldnames=["RMU", "LINK"])
                w.writeheader(); w.writerow({"RMU":"R1", "LINK":"FALSE"})
            store = ProjectStore(root / "project")
            try:
                store.set_source("zenon_db", zdb_path)
                store.set_source("adms_db", adb_path)
                store.set_source("adms_sld", asld_path)
                rows, _ = build_comparison(store)
                row = next(r for r in rows if r["rmu"] == "R1")
                self.assertEqual(row["analysis_ip"], "FALSE")
                self.assertEqual(row["analysis_link"], "FALSE")
                self.assertIn("Driver info: 172.16.1.10", row["analysis_ip_detail"])
                self.assertIn("ADMS SLD LINK: FALSE", row["analysis_link_detail"])
                self.assertIn("IP mismatch", row["remarks"])
                self.assertIn("RMU not linked", row["remarks"])
            finally:
                store.db.close()

    def test_all_blank_analysis_is_blank_but_single_value_is_true(self):
        import csv
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            se_path = root / "SE.xlsx"
            wb = Workbook(); ws = wb.active
            ws.append(["SS", "FEEDER", "EQUIPMENT", "EQUIP. TYPE", "OH / UG"])
            ws.append(["ADF", "", "R1", "", "UG"])
            wb.save(se_path)
            wb.close()
            store = ProjectStore(root / "project")
            try:
                store.set_source("se_list", se_path)
                rows, _ = build_comparison(store)
                row = rows[0]
                self.assertEqual(row["analysis_name"], "TRUE")   # one NAME source
                self.assertEqual(row["analysis_feeder"], "")    # all feeder sources blank
                self.assertEqual(row["analysis_smart"], "")     # all SMART sources blank
                self.assertEqual(row["analysis_type"], "")      # all TYPE sources blank
                self.assertEqual(row["analysis_ip"], "")        # both IP sources blank
                self.assertEqual(row["analysis_link"], "")      # ADMS SLD row missing -> N/A
            finally:
                store.db.close()


class ComparisonReviewSchemaTests(unittest.TestCase):
    def test_comparison_review_columns_are_front_loaded(self):
        from migration_report_tool.schema import COMPARISON_COLUMNS, COLUMNS
        display_keys = [key for key, _label, _width in COMPARISON_COLUMNS]
        report_keys = [key for key, _label, _width in COLUMNS]
        self.assertEqual(
            display_keys[:10],
            ["analysis_name", "analysis_feeder", "analysis_smart", "analysis_type", "analysis_ip", "analysis_link", "remarks", "comments", "no", "rmu"],
        )
        self.assertEqual(report_keys[-8:], ["analysis_name", "analysis_feeder", "analysis_smart", "analysis_type", "analysis_ip", "analysis_link", "remarks", "comments"])
        self.assertEqual(set(display_keys), set(report_keys))
