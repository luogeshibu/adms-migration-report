"""Non-UI validation of comparison, SQLite, template-fill and export behavior."""
from pathlib import Path
import os
import shutil
import tempfile

from openpyxl import load_workbook

from migration_report_tool.core import COLUMNS, ProjectStore, REPORT_MERGES, build_comparison, export_report, workspace_root
from migration_report_tool.domain.mapping.signal_mapping import build_signal_mapping_report_from_store


def run_selftest():
    # Formal regression/self-test must never read or modify the operator's
    # application-global Display Name database.  A real workstation can have
    # legitimate global overrides (for example Device -> RMU Type), which must
    # not change a build test's expected workbook contract.
    old_user_data = os.environ.get("MIGRATION_REPORT_TOOL_USER_DATA_ROOT")
    isolated_user_data = tempfile.TemporaryDirectory()
    os.environ["MIGRATION_REPORT_TOOL_USER_DATA_ROOT"] = str(Path(isolated_user_data.name) / "user-data")
    try:
        _run_selftest_isolated()
    finally:
        if old_user_data is None:
            os.environ.pop("MIGRATION_REPORT_TOOL_USER_DATA_ROOT", None)
        else:
            os.environ["MIGRATION_REPORT_TOOL_USER_DATA_ROOT"] = old_user_data
        isolated_user_data.cleanup()


def _run_selftest_isolated():
    mapping = {
        "se_list": "SE.xlsx",
        "zenon_db": "ZENON-DB.csv",
        "zenon_sld": "ZENON-SLD.csv",
        "adms_db": "ADMS-DB.csv",
        "adms_sld": "ADMS-SLD.csv",
        "ioa": "ZENON-ADMS-IOA.csv",
    }
    root = Path(__file__).resolve().parents[1] / "examples" / "sample-data"
    project = workspace_root() / "_SELFTEST_V080"
    if project.exists():
        shutil.rmtree(project)
    store = ProjectStore(project)
    try:
        for source, filename in mapping.items():
            store.set_source(source, root / filename)
        rows, summary = build_comparison(store)
        assert rows, "comparison returned no rows"
        assert len(COLUMNS) == 56, f"Equipment review RMU profile schema must contain 56 source/review columns, got {len(COLUMNS)}"
        assert "channel_analysis" not in rows[0], "obsolete ADMS Channel Analysis field must not be computed"
        store.save_comparison(rows)
        # Structured Resolution replaces free-form RMU Comments. Pick one issue
        # row and resolve every active FALSE field with the first available source.
        issue_row = next(row for row in rows if any(str(row.get(f"analysis_{field.lower()}", "")).upper() == "FALSE" for field in ("NAME", "FEEDER", "SMART", "TYPE", "IP")))
        issue_rmu = issue_row["rmu"]
        for field in ("NAME", "FEEDER", "SMART", "TYPE", "IP"):
            if str(issue_row.get(f"analysis_{field.lower()}", "")).upper() != "FALSE":
                continue
            candidates = (issue_row.get("resolution_candidates") or {}).get(field, [])
            if candidates:
                candidate = candidates[0]
                store.set_rmu_resolution(
                    issue_rmu, field, "USE_SOURCE", "selftest",
                    selected_source=candidate.get("source", ""), selected_value=candidate.get("value", ""),
                    normalized_value=candidate.get("normalized", ""),
                    analysis_fingerprint=store._rmu_field_fingerprint(issue_row, field),
                )
            else:
                store.set_rmu_resolution(
                    issue_rmu, field, "ACCEPT_EXCEPTION", "selftest",
                    analysis_fingerprint=store._rmu_field_fingerprint(issue_row, field),
                )
        store.sync_rmu_review_from_resolutions(issue_rmu, issue_row, "selftest")
        # Pass RMUs default to Closed in the three-state Review workflow; an explicit
        # Closed decision and comment must remain persisted/exported without changing Analysis.
        manual_pass_row = next(
            row for row in rows
            if row["rmu"] != rows[0]["rmu"]
            and all(str(row.get(f"analysis_{field}", "")).upper() != "FALSE" for field in ("name", "feeder", "smart", "type", "ip"))
        )
        store.update_rmu_review_status(
            manual_pass_row["rmu"], "CLOSED", "selftest",
            reason="Explicit RMU Closed review on automatic Pass result",
        )
        manual_pass_comment = "Optional site verification completed."
        store.update_rmu_manual_review_comment(manual_pass_row["rmu"], manual_pass_comment, "selftest")
        store.update_rmu_check_passed(manual_pass_row["rmu"], True, "selftest")

        # Signal Mapping has the same persistent manual Checked workflow.
        signal_report = build_signal_mapping_report_from_store(store)
        manual_signal = signal_report.rows[0]
        store.update_db_smart_check_passed(
            manual_signal.row_key, manual_signal.rmu, True, "selftest",
            source_hash=signal_report.source_hash, row_hash="selftest-signal-hash",
            site_name=store.folder.name,
        )
        target = export_report(store)
        assert target.exists(), "Excel export not created"

        wb = load_workbook(target, read_only=False, data_only=False)
        try:
            expected_sheets = ["RMU Data Review", "Signal Mapping Review", "STANDARD", "Import Sources", "Change Audit Log"]
            assert wb.sheetnames == expected_sheets, f"unexpected export sheets: {wb.sheetnames}"

            # RMU Data Review must mirror the App review workspace, not the source DATA order.
            ws = wb["RMU Data Review"]
            assert ws["A1"].value == "Review"
            assert ws["A2"].value == "Checked"
            assert ws["B2"].value == "Review"
            assert ws["B3"].value == "CLOSED"
            assert ws["B3"].fill.fgColor.rgb.endswith("2E7D32"), "Automatic pass must use Closed green review color"
            assert ws["C1"].value == "Analysis"
            assert [ws.cell(2, c).value for c in range(3, 8)] == ["NAME", "FEEDER", "SMART", "TYPE", "IP"]
            assert ws["H1"].value == "Remarks"
            assert ws["I1"].value == "Resolution"
            assert ws["J1"].value == "Index"
            assert ws["J2"].value == "No."
            assert ws["K2"].value == "RMU"
            assert ws["L1"].value == "SE"
            assert str(ws["K3"].value) == str(rows[0]["rmu"])
            # Structured issue decisions must be exported in the Resolution column.
            issue_excel_row = next(r for r in range(3, ws.max_row + 1) if str(ws.cell(r, 11).value or "") == str(issue_rmu))
            assert ws.cell(issue_excel_row, 9).value, "Resolution summary must be exported"
            assert ws.cell(issue_excel_row, 2).value == "CLOSED"
            assert ws.cell(issue_excel_row, 2).fill.fgColor.rgb.endswith("2E7D32"), "Closed review must use dark green"
            assert ws.cell(issue_excel_row, 2).font.color.type == "rgb" and ws.cell(issue_excel_row, 2).font.color.rgb.endswith("FFFFFF"), "Closed review must use white text"
            manual_pass_excel_row = next(r for r in range(3, ws.max_row + 1) if str(ws.cell(r, 11).value or "") == str(manual_pass_row["rmu"]))
            assert ws.cell(manual_pass_excel_row, 1).value == "✓", "Checked RMU must export as a visual tick, not PASS text"
            assert ws.cell(manual_pass_excel_row, 1).fill.fgColor.rgb.endswith("2E7D32"), "Checked tick must retain the completed green cell"
            assert ws.cell(manual_pass_excel_row, 2).value == "CLOSED", "Explicit Pass RMU Closed state must be exported"
            assert ws.cell(manual_pass_excel_row, 2).fill.fgColor.rgb.endswith("2E7D32"), "Explicit Closed must use the Closed color"
            assert ws.cell(manual_pass_excel_row, 2).font.color.type == "rgb" and ws.cell(manual_pass_excel_row, 2).font.color.rgb.endswith("FFFFFF"), "Explicit Closed must use white text"
            assert ws.cell(manual_pass_excel_row, 9).value == manual_pass_comment, "Optional Pass RMU comment must be exported in Resolution"

            # Signal Mapping Review mirrors RMU manual verification: Checked first,
            # then Review/Comments, followed by RMU/source groups.
            signal = wb["Signal Mapping Review"]
            assert signal["A1"].value == "Review"
            assert signal["A2"].value == "Checked"
            assert signal["B2"].value == "Review"
            assert signal["C2"].value == "Comments"
            assert signal["D1"].value == "RMU"
            assert signal["D2"].value == "RMU"
            assert signal["E2"].value == "Type"
            assert signal["F1"].value == "STANDARD DATABASE I/O list"
            assert signal["A3"].value == "✓", "Checked signal must export as a visual tick"
            assert signal["A3"].fill.fgColor.rgb.endswith("2E7D32"), "Checked signal tick must use the completed green cell"
            assert signal["B3"].value == "CLOSED"
            assert signal["B3"].fill.fgColor.rgb.endswith("2E7D32"), "Matched signal must default to Closed green"
            assert signal["C3"].fill.fgColor.rgb.endswith("FFFFFF"), "Signal Comments must always stay neutral white"

            assert wb["STANDARD"].max_row > 1
            assert wb["Import Sources"]["A1"].value == "Source Type"
            assert wb["Change Audit Log"]["A1"].value == "ID"
            assert [wb["Change Audit Log"].cell(1, c).value for c in range(1, 5)] == ["ID", "Module", "Record", "Field"]
            # STANDARD is copied from the bundled IOA STANDARD workbook; legacy DATA/DB-smart helpers are absent.
            assert "DATA" not in wb.sheetnames
            assert "DB-smsrt report" not in wb.sheetnames
            assert "DB-smart report" not in wb.sheetnames
        finally:
            wb.close()
        print(f"PASS: {len(rows)} rows; {summary}; export={target.name}")
    finally:
        store.db.close()
        shutil.rmtree(project, ignore_errors=True)


if __name__ == "__main__":
    run_selftest()
