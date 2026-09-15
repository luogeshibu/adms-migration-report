from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from migration_report_tool.infrastructure.database.sqlite_store import ProjectStore
from migration_report_tool.infrastructure.export.workbook_exporter import export_report
from migration_report_tool.version import __version__

ROOT = Path(__file__).resolve().parents[1]
UI_SOURCE = (ROOT / "src/migration_report_tool/ui/main_window.py").read_text(encoding="utf-8")
EXPORT_SOURCE = (ROOT / "src/migration_report_tool/infrastructure/export/workbook_exporter.py").read_text(encoding="utf-8")
STORE_SOURCE = (ROOT / "src/migration_report_tool/infrastructure/database/sqlite_store.py").read_text(encoding="utf-8")


class TestV08194OptionalSignalExportAsyncReview(unittest.TestCase):
    def test_version(self):
        self.assertEqual(__version__, "0.8.196")

    def test_equipment_excel_export_does_not_require_signal_mapping_sources(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            store = ProjectStore(root / "site")
            try:
                store.save_comparison([
                    {
                        "rmu": "EQ-001",
                        "no": 1,
                        "analysis_name": "TRUE",
                        "analysis_feeder": "TRUE",
                        "analysis_smart": "TRUE",
                        "analysis_type": "TRUE",
                        "analysis_ip": "TRUE",
                        "status": "Pass",
                        "remarks": "",
                        "comments": "",
                    }
                ])
                target = export_report(store, target_path=root / "equipment-only.xlsx")
                self.assertTrue(target.exists())
                wb = load_workbook(target, read_only=True, data_only=False)
                try:
                    self.assertEqual(
                        wb.sheetnames,
                        [
                            "RMU Data Review",
                            "Signal Mapping Review",
                            "STANDARD",
                            "Import Sources",
                            "Change Audit Log",
                        ],
                    )
                    ws = wb["Signal Mapping Review"]
                    self.assertEqual(ws["A2"].value, "NOT CONFIGURED")
                    self.assertIn("optional source set is not configured", str(ws["B2"].value))
                finally:
                    wb.close()
            finally:
                store.close()

    def test_export_hot_path_preloads_review_and_resolution_maps(self):
        self.assertIn("all_resolutions = store.rmu_resolution_map()", EXPORT_SOURCE)
        self.assertNotIn("store.rmu_resolution_summary(rmu)", EXPORT_SOURCE)
        self.assertNotIn("store.rmu_manual_review_comment(rmu)", EXPORT_SOURCE)

    def test_excel_export_runs_off_gui_thread_and_uses_fast_row_count_gate(self):
        self.assertIn("def _background_excel_export_job", UI_SOURCE)
        self.assertIn("_background_excel_export_job(project_folder, str(target_path))", UI_SOURCE)
        self.assertIn("self.store.comparison_row_count() <= 0", UI_SOURCE)
        self.assertNotIn("QApplication.setOverrideCursor(Qt.WaitCursor)\n        try:\n            target = export_report", UI_SOURCE)

    def test_equipment_check_clicks_are_batched_off_gui_thread(self):
        self.assertIn("def _background_equipment_check_batch_job", UI_SOURCE)
        self.assertIn('"equipment-check-batch"', UI_SOURCE)
        self.assertIn("timer.setInterval(60)", UI_SOURCE)
        self.assertIn("_commit=False", UI_SOURCE)
        self.assertIn("_commit: bool = True", STORE_SOURCE)

    def test_uncommitted_check_batch_can_be_rolled_back(self):
        with tempfile.TemporaryDirectory() as td:
            folder = Path(td) / "site"
            store = ProjectStore(folder)
            try:
                store.update_rmu_check_passed("EQ-001", True, "tester", _commit=False)
                store.db.rollback()
            finally:
                store.close()
            reopened = ProjectStore(folder)
            try:
                self.assertFalse(bool(int(reopened.rmu_review_record("EQ-001").get("check_passed") or 0)))
            finally:
                reopened.close()


if __name__ == "__main__":
    unittest.main()
