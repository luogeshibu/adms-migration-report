"""Calculated Signal Mapping Review engine.

Signal Mapping Review is application-owned data.  It is **not** read from the
legacy DB-smart worksheet anymore.

Inputs:
* ZENON-ADMS-IOA.csv: RMU + Zenon/ADMS signal mapping rows;
* ADMS-SLD.csv: cabinet type for each RMU;
* active application IOA STANDARD.xlsx: STANDARD reference worksheet (user override or bundled default).

STANDARD is the authoritative expected point list for Signal Mapping Review.
The engine first resolves each RMU Type from ADMS SLD, loads the complete
STANDARD point list for that exact type, then compares only ADMS points from
the combined ZENON-ADMS IOA table against those expected rows. ZENON values are
loaded from the same combined table only after STANDARD/ADMS alignment. Extra
ZENON-only rows retain the semantic ADMS-implementation lookup as an information
hint; ZENON never participates in TRUE/FALSE STANDARD validation.
"""
from __future__ import annotations

from dataclasses import dataclass
from hashlib import sha1, sha256
from pathlib import Path
import json
import re
from typing import Iterable

from openpyxl import load_workbook

from ..analysis.consistency import normalize_type
from ...parsers import clean, normalize_key
from ...infrastructure.parsers import read_mapped_rows
from ...utils.paths import standard_reference_path


@dataclass(frozen=True)
class DBSmartRow:
    excel_row: int
    row_key: str
    rmu: str
    values: tuple[str, ...]
    analysis_detail: str = ""
    suggested_comment: str = ""
    # Stable aliases from earlier Signal Mapping layouts.  v0.8.104 uses these
    # to carry forward existing site review decisions when the STANDARD-driven
    # row identity replaces a legacy source-row identity.
    legacy_row_keys: tuple[str, ...] = ()


@dataclass(frozen=True)
class DBSmartReport:
    source_path: Path
    sheet_name: str
    source_hash: str
    group_definitions: tuple[tuple[str, str, tuple[tuple[str, str, int], ...]], ...]
    columns: tuple[tuple[str, str, int], ...]
    rows: tuple[DBSmartRow, ...]
    analysis_column: int | None
    input_paths: tuple[Path, ...] = ()


SIGNAL_MAPPING_GROUPS = (
    ("RMU", "#E5E7EB", (
        ("rmu_no", "RMU", 105),
        ("rmu_type", "Type", 90),
    )),
    ("STANDARD DATABASE I/O list", "#E5E7EB", (
        ("standard_signal_name", "STANDARD signal name", 250),
        ("standard_dot_no", "STANDARD DOT NO", 135),
    )),
    ("ADMS", "#E5E7EB", (
        ("adms_gss_fid", "ADMS GSS-FID", 200),
        ("adms_signal_name", "ADMS signal name", 270),
        ("adms_dot_no", "ADMS DOT NO", 120),
    )),
    ("ZENON", "#E5E7EB", (
        ("zenon_gss_fid", "ZENON GSS-FID", 200),
        ("zenon_signal_name", "ZENON signal name", 260),
        ("zenon_dot_no", "ZENON DOT NO", 120),
    )),
    ("Analysis", "#E5E7EB", (
        ("analysis_adms_standard", "ADMS/STANDARD", 125),
    )),
    ("Comparison Summary", "#E5E7EB", (
        ("summary_item", "Item", 140),
        ("summary_adms_standard", "ADMS/STANDARD", 130),
    )),
)
SIGNAL_MAPPING_COLUMNS = tuple(
    column for _group, _color, columns in SIGNAL_MAPPING_GROUPS for column in columns
)
ANALYSIS_COLUMN = next(
    i for i, (key, _label, _width) in enumerate(SIGNAL_MAPPING_COLUMNS)
    if key == "analysis_adms_standard"
)


def signal_row_is_zenon_only(row: DBSmartRow) -> bool:
    """True when a Signal Mapping row has ZENON data but no ADMS/STANDARD data.

    This is a presentation/review workflow classification only; the calculated
    validation engine and source values remain untouched.
    """
    index = {key: i for i, (key, _label, _width) in enumerate(SIGNAL_MAPPING_COLUMNS)}

    def present(*keys: str) -> bool:
        return any(i < len(row.values) and clean(row.values[i]) for key in keys if (i := index.get(key)) is not None)

    zenon = present("zenon_gss_fid", "zenon_signal_name", "zenon_dot_no")
    adms = present("adms_gss_fid", "adms_signal_name", "adms_dot_no")
    standard = present("standard_signal_name", "standard_dot_no")
    return bool(zenon and not adms and not standard)


def _file_sha256(path: Path, chunk_size: int = 1024 * 1024) -> str:
    digest = sha256()
    with Path(path).open("rb") as handle:
        while True:
            block = handle.read(chunk_size)
            if not block:
                break
            digest.update(block)
    return digest.hexdigest()


def _combined_hash(paths: Iterable[Path]) -> str:
    digest = sha256()
    for path in paths:
        path = Path(path)
        digest.update(path.name.encode("utf-8", errors="replace"))
        digest.update(_file_sha256(path).encode("ascii"))
    return digest.hexdigest()


def _norm_header(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", "", clean(value).casefold())


def _norm_type_key(value: object) -> str:
    # STANDARD Type is a business key, not free text. Keep extended types such
    # as "3L1T-MRMU (...)" distinct from the base "3L1T" so Type+IOA remains
    # a deterministic unique lookup as required by the report contract.
    return re.sub(r"\s+", " ", clean(value).strip().upper())


def _norm_ioa(value: object) -> str:
    text = clean(value)
    if not text:
        return ""
    # Excel/CSV frequently represents integer IOAs as 1000.0.
    try:
        number = float(text.replace(",", ""))
        if number.is_integer():
            return str(int(number))
    except Exception:
        pass
    return text.upper()


def _display_value(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _strip_rmu_prefix(signal: object, rmu: object) -> str:
    text = clean(signal).strip()
    rmu_text = normalize_key(rmu).strip()
    if not text:
        return ""
    if rmu_text:
        # ADMS examples: "26859 Y1 CMD".  Remove only a true leading RMU token.
        text = re.sub(rf"^\s*{re.escape(rmu_text)}(?:\s+|[_-]+)", "", text, flags=re.I)
    return text.strip()


def normalize_signal_name(value: object, *, strip_rmu: object = "") -> str:
    """Normalize signal text using the legacy report's proven comparison rules."""
    text = _strip_rmu_prefix(value, strip_rmu) if strip_rmu else clean(value)
    text = text.casefold().replace(" ", "").replace("_", "")
    replacements = (
        ("iavalue(a)", "icurrent(a)"),
        ("ibvalue(a)", "icurrent(b)"),
        ("icvalue(a)", "icurrent(c)"),
        ("invalue(a)", "icurrent(n)"),
        ("uabvalue(kv)", "vvoltage(a-b)"),
        ("ubcvalue(kv)", "vvoltage(b-c)"),
        ("ucavalue(kv)", "vvoltage(c-a)"),
        ("unvalue(kv)", "vvoltage(n)"),
        ("value", ""),
        ("state", ""),
    )
    for old, new in replacements:
        text = text.replace(old, new)
    return text


def _find_standard_sheet(wb):
    for name in wb.sheetnames:
        if name.strip().casefold() == "standard":
            return wb[name]
    raise ValueError("STANDARD sheet not found in the active IOA STANDARD workbook.")


def _standard_header_columns(ws) -> tuple[int, int, int, int]:
    """Return (header_row, type_col, ioa_col, name_col)."""
    for row_idx in range(1, min(ws.max_row, 25) + 1):
        found: dict[str, int] = {}
        for col_idx in range(1, ws.max_column + 1):
            key = _norm_header(ws.cell(row_idx, col_idx).value)
            if key == "type" and "type" not in found:
                found["type"] = col_idx
            elif key == "ioa" and "ioa" not in found:
                found["ioa"] = col_idx
            elif key == "name" and "name" not in found:
                found["name"] = col_idx
        if {"type", "ioa", "name"} <= set(found):
            return row_idx, found["type"], found["ioa"], found["name"]
    raise ValueError("STANDARD must contain columns named Type, IOA and name.")


def _read_standard_index(path: Path, overrides: dict[str, str] | None = None):
    """Return STANDARD rows indexed by exact RMU Type and Type+IOA.

    STANDARD is authoritative.  Cross-Type fallback is deliberately forbidden:
    a 2L1T RMU can only be validated against the 2L1T STANDARD point list, etc.
    """
    mapped = read_mapped_rows("standard_reference", path, overrides or {}, strict=True)
    by_type: dict[str, list[dict[str, str]]] = {}
    by_type_ioa: dict[tuple[str, str], list[dict[str, str]]] = {}
    for row_number, row in enumerate(mapped.rows, 1):
        type_value = _norm_type_key(row.get("type"))
        ioa_value = _norm_ioa(row.get("ioa"))
        name_value = _display_value(row.get("name"))
        if not type_value or not ioa_value:
            continue
        item = {
            "type": type_value,
            "ioa": ioa_value,
            "name": name_value,
            "row": str(row_number),
        }
        by_type.setdefault(type_value, []).append(item)
        by_type_ioa.setdefault((type_value, ioa_value), []).append(item)
    return by_type, by_type_ioa

def _first_value(row: dict, *keys: str) -> str:
    # First exact lookup, then case/character-insensitive lookup for site variants.
    for key in keys:
        if clean(row.get(key)):
            return clean(row.get(key))
    normalized = {_norm_header(k): clean(v) for k, v in row.items() if clean(v)}
    for key in keys:
        hit = normalized.get(_norm_header(key))
        if hit:
            return hit
    return ""


def _read_adms_sld_types(path: Path, overrides: dict[str, str] | None = None, *, sheet_name: str | None = None):
    rows = read_mapped_rows("adms_sld", path, overrides or {}, strict=True, sheet_name=sheet_name).rows
    grouped: dict[str, list[dict]] = {}
    for row in rows:
        rmu = normalize_key(row.get("rmu"))
        if rmu:
            grouped.setdefault(rmu, []).append(row)

    result: dict[str, tuple[str, str]] = {}
    for rmu, candidates in grouped.items():
        types = {normalize_type(row.get("rmu_type")) for row in candidates}
        types.discard("")
        if len(types) == 1:
            result[rmu] = (next(iter(types)), "")
        elif len(types) > 1:
            result[rmu] = ("", f"ADMS SLD has conflicting RMU types: {', '.join(sorted(types))}")
        else:
            result[rmu] = ("", "ADMS SLD RMU type is blank")
    return result


def _row_key(source: dict, row_number: int) -> str:
    identity = [
        normalize_key(source.get("rmu")),
        clean(source.get("zenon_gss_fid")),
        clean(source.get("zenon_signal_name")),
        _norm_ioa(source.get("zenon_dot_no")),
        clean(source.get("adms_gss_fid")),
        clean(source.get("adms_signal_name")),
        _norm_ioa(source.get("adms_dot_no")),
    ]
    payload = "\x1f".join(x.casefold() for x in identity)
    if not payload.strip("\x1f"):
        payload = f"row:{row_number}"
    return sha1(payload.encode("utf-8", errors="replace")).hexdigest()


def _standard_row_key(rmu: str, rmu_type: str, standard: dict[str, str]) -> str:
    payload = "\x1f".join((
        "standard",
        normalize_key(rmu).casefold(),
        _norm_type_key(rmu_type).casefold(),
        _norm_ioa(standard.get("ioa")).casefold(),
        normalize_signal_name(standard.get("name")).casefold(),
    ))
    return sha1(payload.encode("utf-8", errors="replace")).hexdigest()


def _standard_driven_analysis(
    *, rmu: str, rmu_type: str, type_reason: str,
    standard: dict[str, str] | None, adms_signal: str, adms_dot: str,
) -> tuple[str, str]:
    """Validate one ADMS point against the exact-Type STANDARD expectation."""
    if standard is None:
        detail = [
            "ADMS/STANDARD consistency check",
            "",
            "Rule: the RMU Type selects the authoritative STANDARD point list before ADMS comparison.",
            f"RMU: {rmu or '—'}",
            f"Type (from ADMS SLD): {rmu_type or '<missing>'}",
            f"ADMS_DOT_NO: {adms_dot or '<blank>'}",
            f"ADMS signal: {adms_signal or '<blank>'}",
            "",
            "Result: FALSE",
            "Reason: this ADMS point is not present in the STANDARD point list for the resolved RMU Type.",
        ]
        if type_reason:
            detail.append("Type context: " + type_reason)
        return "FALSE", "\n".join(detail)

    standard_name = clean(standard.get("name"))
    standard_dot = _norm_ioa(standard.get("ioa"))
    adms_norm = normalize_signal_name(adms_signal, strip_rmu=rmu)
    standard_norm = normalize_signal_name(standard_name)
    dot_match = bool(adms_dot and _norm_ioa(adms_dot) == standard_dot)
    name_match = bool(adms_norm and standard_norm and adms_norm == standard_norm)
    result = "TRUE" if dot_match and name_match else "FALSE"
    detail = [
        "ADMS/STANDARD consistency check",
        "",
        "Rule: RMU Type selects the STANDARD point list; within that type, normalized name + point number must match.",
        f"RMU: {rmu or '—'}",
        f"Type (from ADMS SLD): {rmu_type or '<missing>'}",
        f"STANDARD Type: {clean(standard.get('type')) or '<blank>'}",
        f"STANDARD IOA: {standard_dot}",
        f"STANDARD name: {standard_name or '<blank>'}",
        f"ADMS_DOT_NO: {adms_dot or '<missing>'}",
        f"ADMS signal: {adms_signal or '<missing>'}",
        "",
        f"Normalized STANDARD name: {standard_norm or '<blank>'}",
        f"Normalized ADMS signal: {adms_norm or '<blank>'}",
        "",
        f"Result: {result}",
    ]
    if result == "TRUE":
        detail.append("Reason: ADMS point matches the STANDARD expectation for this RMU Type.")
    elif not clean(adms_signal) and not adms_dot:
        detail.append("Reason: STANDARD requires this point, but ADMS has no corresponding point.")
    else:
        problems = []
        if not dot_match:
            problems.append("point number mismatch")
        if not name_match:
            problems.append("signal name mismatch")
        detail.append("Reason: " + " / ".join(problems or ["STANDARD mismatch"]))
    if type_reason:
        detail.append("Type context: " + type_reason)
    return result, "\n".join(detail)


def _analysis_for_row(
    *, rmu: str, rmu_type: str, type_reason: str, adms_signal: str,
    adms_dot: str, type_candidates: list[dict[str, str]],
    ioa_candidates: list[dict[str, str]],
) -> tuple[str, str, str, str]:
    """Return analysis, STANDARD name/DOT and human-readable reason.

    Business rule:
      ADMS signal name + ADMS DOT number == STANDARD name + IOA  -> TRUE

    RMU Type is supporting context only. It is used to choose the most useful
    STANDARD row for mismatch explanation, but a Type mismatch must not turn an
    otherwise exact name+DOT match into FALSE.
    """
    if not adms_dot and not clean(adms_signal):
        return "", "", "", "ADMS signal and ADMS point number are blank; no STANDARD check was performed."
    if not adms_dot:
        return "FALSE", "", "", f"ADMS_DOT_NO is blank for RMU {rmu or '—'}; signal name + DOT are required for STANDARD validation."

    adms_norm = normalize_signal_name(adms_signal, strip_rmu=rmu)
    ioa_norm = _norm_ioa(adms_dot)

    # Primary rule: name + DOT/IOA equality is sufficient, across RMU types.
    exact_value_matches = [
        item for item in ioa_candidates
        if adms_norm
        and normalize_signal_name(item.get("name"))
        and normalize_signal_name(item.get("name")) == adms_norm
        and _norm_ioa(item.get("ioa")) == ioa_norm
    ]
    if exact_value_matches:
        # Prefer the same Type for presentation when one exists, but do not
        # require Type equality for the TRUE result.
        same_type = [item for item in exact_value_matches if rmu_type and item.get("type") == _norm_type_key(rmu_type)]
        standard = (same_type or exact_value_matches)[0]
        standard_name = clean(standard.get("name"))
        standard_dot = _norm_ioa(standard.get("ioa"))
        detail = [
            "ADMS/STANDARD consistency check",
            "",
            "Rule: normalized ADMS signal name + ADMS DOT number must match STANDARD name + IOA.",
            "RMU Type is supporting context and is not required when name + DOT already match.",
            "",
            f"RMU: {rmu or '—'}",
            f"Type (from ADMS SLD): {rmu_type or '<missing>'}",
            f"STANDARD Type: {standard.get('type') or '<blank>'}",
            f"ADMS_DOT_NO: {adms_dot}",
            f"STANDARD IOA: {standard_dot}",
            f"ADMS signal: {adms_signal or '<blank>'}",
            f"STANDARD name: {standard_name}",
            "",
            f"Normalized ADMS signal: {adms_norm or '<blank>'}",
            f"Normalized STANDARD name: {normalize_signal_name(standard_name) or '<blank>'}",
            "",
            "Result: TRUE",
            "Reason: ADMS signal name and DOT number match the STANDARD reference.",
        ]
        return "TRUE", standard_name, standard_dot, "\n".join(detail)

    # No exact name+DOT match. Use Type+IOA (when available) to show the
    # expected STANDARD row and explain the mismatch. If Type is unavailable or
    # wrong, a unique IOA row can still provide useful context.
    candidates = list(type_candidates)
    lookup_reason = "Type + IOA"
    if not candidates:
        candidates = list(ioa_candidates)
        lookup_reason = "IOA"

    if not candidates:
        type_text = rmu_type or "<missing>"
        return "FALSE", "", "", (
            f"STANDARD mapping not found for ADMS signal={adms_signal or '<blank>'}, "
            f"IOA={adms_dot}, Type={type_text}. No STANDARD row has the same DOT number."
        )

    # If several candidates share the IOA, choose a same-Type candidate first
    # for display. Multiple rows are not an automatic failure anymore because
    # the business truth condition is name+DOT; at this point none matched the
    # ADMS signal name, so the result is FALSE regardless.
    same_type = [item for item in candidates if rmu_type and item.get("type") == _norm_type_key(rmu_type)]
    standard = (same_type or candidates)[0]
    standard_name = clean(standard.get("name"))
    standard_dot = _norm_ioa(standard.get("ioa"))
    standard_norm = normalize_signal_name(standard_name)
    dot_match = ioa_norm == standard_dot
    name_match = bool(adms_norm and standard_norm and adms_norm == standard_norm)

    detail = [
        "ADMS/STANDARD consistency check",
        "",
        "Rule: normalized ADMS signal name + ADMS DOT number must match STANDARD name + IOA.",
        "RMU Type is supporting context only.",
        "",
        f"RMU: {rmu or '—'}",
        f"Type (from ADMS SLD): {rmu_type or '<missing>'}",
        f"STANDARD Type: {standard.get('type') or '<blank>'}",
        f"Fallback lookup: {lookup_reason}",
        f"ADMS_DOT_NO: {adms_dot}",
        f"STANDARD IOA: {standard_dot}",
        f"ADMS signal: {adms_signal or '<blank>'}",
        f"STANDARD name: {standard_name or '<blank>'}",
        "",
        f"Normalized ADMS signal: {adms_norm or '<blank>'}",
        f"Normalized STANDARD name: {standard_norm or '<blank>'}",
        "",
        "Result: FALSE",
    ]
    problems = []
    if not name_match:
        problems.append("signal name mismatch")
    if not dot_match:
        problems.append("point number mismatch")
    if not problems:
        # Defensive only: exact name+DOT would already have returned TRUE.
        problems.append("STANDARD value mismatch")
    detail.append("Reason: " + " / ".join(problems))
    if type_reason:
        detail.append("Type context: " + type_reason)
    return "FALSE", standard_name, standard_dot, "\n".join(detail)


def _signal_equipment_token(value: object, *, zenon: bool = False) -> tuple[str, str]:
    """Return the business equipment token plus the human matching explanation.

    ZENON uses ``TR1`` for the transformer-side channel while ADMS uses ``Q1``.
    This translation is deliberately limited to matching the implementation
    location for comments; it never rewrites source values or validation data.
    """
    text = clean(value).upper()
    match = re.search(r"(?:^|[^A-Z0-9])(TR|Y|Q)(\d+)(?=$|[^A-Z0-9])", text)
    if not match:
        return "", ""
    family, number = match.groups()
    raw = f"{family}{number}"
    if zenon and family == "TR":
        return f"Q{number}", f"{raw}→Q{number}"
    return raw, raw


def _signal_measurement_token(value: object, *, zenon: bool = False) -> str:
    """Normalize a signal measurement into a small semantic vocabulary.

    The goal is intentionally conservative: identify the signal that already
    exists somewhere else in ADMS for a ZENON-only row.  No match produced by
    this helper changes validation or Review status.
    """
    text = clean(value).upper().replace(" ", "_")
    if not text:
        return ""

    # Order matters (KVAR before KVA, VOLT_N before generic VOLT, etc.).
    if "KVAR" in text or "Q(KVAR)" in text.replace("_", ""):
        return "KVAR"
    if "KVA" in text or "S(KVA)" in text.replace("_", ""):
        return "KVA"
    if re.search(r"(?:^|[_\-])KW(?:$|[_\-])", text) or "P(KW)" in text.replace("_", ""):
        return "KW"
    if "POWER_FACTOR" in text or re.search(r"(?:^|[_\-])PF(?:$|[_\-])", text):
        return "PF"

    if "AMP_A" in text or "I_A_VALUE" in text or "I_CURRENT(A)" in text.replace("_", ""):
        return "AMP_A"
    if "AMP_B" in text or "I_B_VALUE" in text or "I_CURRENT(B)" in text.replace("_", ""):
        return "AMP_B"
    if "AMP_C" in text or "I_C_VALUE" in text or "I_CURRENT(C)" in text.replace("_", ""):
        return "AMP_C"
    if "AMP_N" in text or "I_N_VALUE" in text or "I_CURRENT(N)" in text.replace("_", ""):
        return "AMP_N"

    compact = text.replace("_", "")
    if "VOLT_N" in text or re.search(r"(?:^|[_\-])VN(?:$|[_\-])", text) or "U_N_VALUE" in text or "VVOLTAGE(N)" in compact:
        return "VOLT_N"
    if "VOLT_A" in text or "U_AB_VALUE" in text or "VVOLTAGE(A-B)" in compact:
        return "VOLT_A"
    if "VOLT_B" in text or "U_BC_VALUE" in text or "VVOLTAGE(B-C)" in compact:
        return "VOLT_B"
    if "VOLT_C" in text or "U_CA_VALUE" in text or "VVOLTAGE(C-A)" in compact:
        return "VOLT_C"

    if "ST_LOCK" in text or "LOCK/UNLOCK" in clean(value).upper():
        return "ST_LOCK"
    if "ST_GND" in text:
        return "ST_GND"
    if re.search(r"(?:^|[_\-])CMD(?:$|[_\-])", text):
        return "CMD"
    if re.search(r"(?:^|[_\-])ST(?:$|[_\-])", text) or clean(value).upper().rstrip().endswith(" STATE"):
        return "ST"
    return ""


def _adms_candidate_signature(signal_name: object) -> tuple[str, str]:
    equipment, _ = _signal_equipment_token(signal_name, zenon=False)
    measurement = _signal_measurement_token(signal_name, zenon=False)
    # Ground-state ADMS names are usually KQ1/KY1 state.  Detect the channel
    # from the embedded Q/Y token while classifying the measurement as ST_GND.
    text = clean(signal_name).upper()
    ground = re.search(r"(?:^|[^A-Z0-9])K(Q|Y)(\d+)\s*STATE", text)
    if ground:
        equipment = f"{ground.group(1)}{ground.group(2)}"
        measurement = "ST_GND"
    return equipment, measurement


def _build_zenon_only_adms_comment(source: dict, adms_rows_for_rmu: list[dict]) -> str:
    """Describe where a ZENON-only signal is implemented in ADMS.

    Matching is semantic and intentionally ignores DOT alignment.  This is
    important for legacy rows such as Y1 KVAR, where the ZENON point may sit on
    13057 while the ADMS implementation is Y1 Q(kVar) on 13059.  ZENON TRn is
    matched to ADMS Qn as required by the project naming convention.
    """
    zenon_signal = clean(source.get("zenon_signal_name"))
    if not zenon_signal:
        return ""
    target_equipment, equipment_note = _signal_equipment_token(zenon_signal, zenon=True)
    target_measurement = _signal_measurement_token(zenon_signal, zenon=True)
    if not target_equipment or not target_measurement:
        return (
            f"Automatic ADMS implementation lookup could not derive a unique keyword signature from "
            f"ZENON signal '{zenon_signal}'."
        )

    matches: list[dict] = []
    for candidate in adms_rows_for_rmu:
        adms_signal = clean(candidate.get("adms_signal_name"))
        if not adms_signal:
            continue
        equipment, measurement = _adms_candidate_signature(adms_signal)
        if equipment == target_equipment and measurement == target_measurement:
            matches.append(candidate)

    match_reason = f"{equipment_note or target_equipment} + {target_measurement}"
    if not matches:
        return (
            f"ADMS implementation not found automatically for ZENON signal '{zenon_signal}' "
            f"(searched by {match_reason})."
        )

    # Deduplicate identical ADMS locations that can appear more than once in a
    # source file.  Keep deterministic source order for the customer-facing text.
    unique: list[tuple[str, str, str]] = []
    seen = set()
    for candidate in matches:
        item = (
            clean(candidate.get("adms_gss_fid")),
            clean(candidate.get("adms_signal_name")),
            _norm_ioa(candidate.get("adms_dot_no")),
        )
        if item in seen:
            continue
        seen.add(item)
        unique.append(item)

    rendered = []
    for gss, signal, dot in unique[:3]:
        parts = [signal]
        if dot:
            parts.append(f"ADMS DOT {dot}")
        if gss:
            parts.append(f"GSS-FID {gss}")
        rendered.append(" | ".join(parts))
    prefix = "ADMS implementation" if len(unique) == 1 else "ADMS implementations"
    suffix = f" (matched by {match_reason})."
    if len(unique) > 3:
        suffix = f"; +{len(unique) - 3} additional match(es)" + suffix
    return f"{prefix}: " + "; ".join(rendered) + suffix



def zenon_only_adms_match_rows(source_row: DBSmartRow, rows: Iterable[DBSmartRow]) -> tuple[DBSmartRow, ...]:
    """Return semantic ADMS implementation row(s) for one ZENON-only row.

    This exposes the same semantic lookup used to generate the automatic
    customer-facing comment, but returns stable report rows so the UI can
    highlight and navigate to the actual ADMS implementation.  Matching is
    deliberately independent of DOT alignment and applies the project TRn→Qn
    convention for lookup only.
    """
    if not signal_row_is_zenon_only(source_row):
        return ()
    index = {key: i for i, (key, _label, _width) in enumerate(SIGNAL_MAPPING_COLUMNS)}
    zenon_idx = index.get("zenon_signal_name")
    if zenon_idx is None or zenon_idx >= len(source_row.values):
        return ()
    zenon_signal = clean(source_row.values[zenon_idx])
    target_equipment, _equipment_note = _signal_equipment_token(zenon_signal, zenon=True)
    target_measurement = _signal_measurement_token(zenon_signal, zenon=True)
    if not target_equipment or not target_measurement:
        return ()

    adms_idx = index.get("adms_signal_name")
    matched: list[DBSmartRow] = []
    seen: set[tuple[str, str, str]] = set()
    for candidate in rows:
        if clean(candidate.rmu) != clean(source_row.rmu):
            continue
        if adms_idx is None or adms_idx >= len(candidate.values):
            continue
        adms_signal = clean(candidate.values[adms_idx])
        if not adms_signal:
            continue
        equipment, measurement = _adms_candidate_signature(adms_signal)
        if equipment != target_equipment or measurement != target_measurement:
            continue
        gss = clean(candidate.values[index["adms_gss_fid"]]) if index.get("adms_gss_fid") is not None else ""
        dot = clean(candidate.values[index["adms_dot_no"]]) if index.get("adms_dot_no") is not None else ""
        signature = (gss, adms_signal, _norm_ioa(dot))
        if signature in seen:
            continue
        seen.add(signature)
        matched.append(candidate)
    return tuple(matched)

def _legacy_source_row_keys(ioa_rows: list[dict]) -> dict[int, str]:
    """Reproduce the pre-v0.8.99 source-row identities for review migration."""
    seen: dict[str, int] = {}
    result: dict[int, str] = {}
    for source_row_number, source in enumerate(ioa_rows, 2):
        base_key = _row_key(source, source_row_number)
        occurrence = seen.get(base_key, 0)
        seen[base_key] = occurrence + 1
        result[source_row_number] = base_key if occurrence == 0 else f"{base_key}:{occurrence + 1}"
    return result


def signal_review_alias_map(report: DBSmartReport) -> dict[str, str]:
    """Map historical Signal review row keys to the current canonical row key."""
    aliases: dict[str, str] = {}
    for row in report.rows:
        for legacy in getattr(row, "legacy_row_keys", ()) or ():
            legacy = clean(legacy)
            if legacy and legacy != row.row_key:
                aliases[legacy] = row.row_key
    return aliases


def _point_number(value: object) -> int | None:
    text = clean(value)
    if not text:
        return None
    try:
        return int(float(text))
    except (TypeError, ValueError):
        match = re.search(r"-?\d+(?:\.\d+)?", text)
        if not match:
            return None
        try:
            return int(float(match.group(0)))
        except ValueError:
            return None


def signal_review_metadata(report: DBSmartReport, row: DBSmartRow) -> dict[str, str]:
    """Return persistent business metadata for one Signal review row.

    This metadata is deliberately independent from the current row_key so an
    unresolved NEEDS ACTION remains classifiable at report-export time even if
    a later source refresh changes the calculated row identity.
    """
    index = {key: i for i, (key, _label, _width) in enumerate(report.columns)}

    def value(key: str) -> str:
        i = index.get(key)
        return clean(row.values[i]) if i is not None and i < len(row.values) else ""

    point_text = value("adms_dot_no") or value("standard_dot_no") or value("zenon_dot_no")
    point_no = _point_number(point_text)
    category = "ANALOG" if point_no is not None and point_no >= 13000 else "STATUS_CMD"
    signal_name = value("adms_signal_name") or value("standard_signal_name") or value("zenon_signal_name")
    snapshot = {
        "rmu": clean(row.rmu),
        "rmu_type": value("rmu_type"),
        "standard_signal": value("standard_signal_name"),
        "standard_dot": value("standard_dot_no"),
        "adms_signal": value("adms_signal_name"),
        "adms_dot": value("adms_dot_no"),
        "zenon_signal": value("zenon_signal_name"),
        "zenon_dot": value("zenon_dot_no"),
        "analysis": clean(row.values[report.analysis_column]) if report.analysis_column is not None and report.analysis_column < len(row.values) else "",
        "analysis_detail": clean(getattr(row, "analysis_detail", "")),
        "suggested_comment": clean(getattr(row, "suggested_comment", "")),
    }
    return {
        "signal_category": category,
        "point_no": str(point_no) if point_no is not None else point_text,
        "signal_name": signal_name,
        "signal_snapshot_json": json.dumps(snapshot, ensure_ascii=False),
    }


def build_signal_mapping_report(
    ioa_path: Path, adms_sld_path: Path, standard_workbook_path: Path, *,
    ioa_overrides: dict[str, str] | None = None,
    adms_sld_overrides: dict[str, str] | None = None,
    standard_overrides: dict[str, str] | None = None,
    ioa_sheet_name: str | None = None,
    adms_sld_sheet_name: str | None = None,
) -> DBSmartReport:
    """Build Signal Mapping in STANDARD -> ADMS -> ZENON-hint order.

    Business contract:
    1. Resolve RMU Type from ADMS SLD.
    2. Select the exact-Type STANDARD point list.
    3. Compare STANDARD against the ADMS half of ZENON-ADMS-IOA.csv.
       Only this comparison produces TRUE/FALSE.
    4. Load/show the ZENON half of the same combined table.  A ZENON-only
       point may produce an ADMS implementation hint, but never changes the
       STANDARD/ADMS TRUE/FALSE result or validation totals.
    """
    ioa_path = Path(ioa_path)
    adms_sld_path = Path(adms_sld_path)
    standard_workbook_path = Path(standard_workbook_path)
    for path, label in (
        (standard_workbook_path, "active IOA STANDARD.xlsx"),
        (adms_sld_path, "ADMS-SLD.csv"),
        (ioa_path, "ZENON-ADMS-IOA.csv"),
    ):
        if not path.exists():
            raise FileNotFoundError(f"{label} not found: {path}")

    # Resolve cabinet Type and STANDARD first.  ADMS and ZENON are then read
    # from the two halves of the same combined IOA source table.
    adms_types = _read_adms_sld_types(adms_sld_path, adms_sld_overrides, sheet_name=adms_sld_sheet_name)
    standard_by_type, _standard_type_ioa = _read_standard_index(standard_workbook_path, standard_overrides)
    ioa_rows = read_mapped_rows("ioa", ioa_path, ioa_overrides or {}, strict=True, sheet_name=ioa_sheet_name).rows
    legacy_keys = _legacy_source_row_keys(ioa_rows)

    ioa_by_rmu: dict[str, list[tuple[int, dict, str]]] = {}
    adms_rows_by_rmu: dict[str, list[dict]] = {}
    for source_row_number, source in enumerate(ioa_rows, 2):
        rmu = normalize_key(source.get("rmu"))
        if not rmu:
            continue
        legacy_key = legacy_keys[source_row_number]
        ioa_by_rmu.setdefault(rmu, []).append((source_row_number, source, legacy_key))
        if clean(source.get("adms_signal_name")) or clean(source.get("adms_dot_no")) or clean(source.get("adms_gss_fid")):
            adms_rows_by_rmu.setdefault(rmu, []).append(source)

    all_rmus = sorted(set(ioa_by_rmu), key=lambda value: (0, int(value)) if value.isdigit() else (1, value.casefold()))
    detail_rows: list[dict] = []
    adms_results: list[bool] = []

    def source_values(source: dict | None):
        source = source or {}
        return (
            clean(source.get("adms_gss_fid")), clean(source.get("adms_signal_name")), _norm_ioa(source.get("adms_dot_no")),
            clean(source.get("zenon_gss_fid")), clean(source.get("zenon_signal_name")), _norm_ioa(source.get("zenon_dot_no")),
        )

    def append_row(*, rmu: str, rmu_type: str, standard: dict | None, source_no: int, source: dict | None,
                   analysis: str, detail: str, row_key: str, suggested_comment: str = "", legacy_aliases: tuple[str, ...] = ()):
        adms_gss, adms_signal, adms_dot, zenon_gss, zenon_signal, zenon_dot = source_values(source)
        standard_name = clean((standard or {}).get("name"))
        standard_dot = _norm_ioa((standard or {}).get("ioa"))
        detail_rows.append({
            "source_row_number": source_no,
            "row_key": row_key,
            "legacy_row_keys": tuple(key for key in legacy_aliases if clean(key) and clean(key) != row_key),
            "rmu": rmu,
            "analysis_detail": detail,
            "suggested_comment": suggested_comment,
            "values": [
                rmu, rmu_type,
                standard_name, standard_dot,
                adms_gss, adms_signal, adms_dot,
                zenon_gss, zenon_signal, zenon_dot,
                analysis,
            ],
        })
        if analysis in {"TRUE", "FALSE"}:
            adms_results.append(analysis == "TRUE")

    for rmu in all_rmus:
        rmu_type, type_reason = adms_types.get(rmu, ("", f"RMU {rmu or '—'} was not found in ADMS SLD."))
        type_key = _norm_type_key(rmu_type)
        standard_rows = list(standard_by_type.get(type_key, ())) if type_key else []
        source_rows = list(ioa_by_rmu.get(rmu, ()))
        used_source_indexes: set[int] = set()

        # 1) STANDARD-driven expected rows: only ADMS values participate in validation.
        for standard in standard_rows:
            standard_dot = _norm_ioa(standard.get("ioa"))
            candidates = [
                (idx, source_no, source, legacy_key)
                for idx, (source_no, source, legacy_key) in enumerate(source_rows)
                if idx not in used_source_indexes and _norm_ioa(source.get("adms_dot_no")) == standard_dot
                and (clean(source.get("adms_signal_name")) or clean(source.get("adms_gss_fid")) or clean(source.get("adms_dot_no")))
            ]
            standard_norm = normalize_signal_name(standard.get("name"))
            exact = [
                item for item in candidates
                if normalize_signal_name(item[2].get("adms_signal_name"), strip_rmu=rmu) == standard_norm
            ]
            chosen = (exact or candidates)[0] if (exact or candidates) else None
            source_idx = chosen[0] if chosen else None
            source_no = chosen[1] if chosen else 0
            source = chosen[2] if chosen else None
            legacy_key = chosen[3] if chosen else ""
            if source_idx is not None:
                used_source_indexes.add(source_idx)
            _ag, adms_signal, adms_dot, _zg, _zs, _zd = source_values(source)
            analysis, detail = _standard_driven_analysis(
                rmu=rmu, rmu_type=rmu_type, type_reason=type_reason,
                standard=standard, adms_signal=adms_signal, adms_dot=adms_dot,
            )
            append_row(
                rmu=rmu, rmu_type=rmu_type, standard=standard, source_no=source_no, source=source,
                analysis=analysis, detail=detail, row_key=_standard_row_key(rmu, rmu_type, standard),
                legacy_aliases=((legacy_key,) if legacy_key else ()),
            )

        # 2) Any remaining ADMS point is extra for this RMU Type -> FALSE.
        for idx, (source_no, source, legacy_key) in enumerate(source_rows):
            if idx in used_source_indexes:
                continue
            adms_present = bool(clean(source.get("adms_signal_name")) or clean(source.get("adms_dot_no")) or clean(source.get("adms_gss_fid")))
            if not adms_present:
                continue
            used_source_indexes.add(idx)
            _ag, adms_signal, adms_dot, _zg, _zs, _zd = source_values(source)
            analysis, detail = _standard_driven_analysis(
                rmu=rmu, rmu_type=rmu_type, type_reason=type_reason,
                standard=None, adms_signal=adms_signal, adms_dot=adms_dot,
            )
            append_row(
                rmu=rmu, rmu_type=rmu_type, standard=None, source_no=source_no, source=source,
                analysis=analysis, detail=detail, row_key=legacy_key,
            )

        # 3) Remaining ZENON-only rows are informational implementation hints only.
        for idx, (source_no, source, legacy_key) in enumerate(source_rows):
            if idx in used_source_indexes:
                continue
            zenon_present = bool(clean(source.get("zenon_signal_name")) or clean(source.get("zenon_dot_no")) or clean(source.get("zenon_gss_fid")))
            if not zenon_present:
                continue
            suggested_comment = _build_zenon_only_adms_comment(source, adms_rows_by_rmu.get(rmu, []))
            detail = (
                "ZENON implementation hint\n\n"
                "This ZENON point is outside the STANDARD/ADMS validation pair. "
                "It does not participate in TRUE/FALSE comparison.\n\n"
                + (suggested_comment or "ADMS implementation could not be determined automatically.")
            )
            append_row(
                rmu=rmu, rmu_type=rmu_type, standard=None, source_no=source_no, source=source,
                analysis="", detail=detail, row_key=legacy_key, suggested_comment=suggested_comment,
            )

    def summary(results: list[bool]):
        matched = sum(1 for x in results if x)
        mismatched = sum(1 for x in results if not x)
        total = len(results)
        return [matched, mismatched, total, (matched / total if total else None), (mismatched / total if total else None)]

    adms_summary = summary(adms_results)
    items = ["Matched", "Mismatched", "Total Compared", "Match Rate", "Mismatch Rate"]

    rows: list[DBSmartRow] = []
    for index, item in enumerate(detail_rows):
        values = list(item["values"])
        if index < len(items):
            def fmt(v):
                if v is None:
                    return ""
                if isinstance(v, float):
                    return f"{v:.2%}"
                return str(v)
            values.extend([items[index], fmt(adms_summary[index])])
        else:
            values.extend(["", ""])
        rows.append(DBSmartRow(
            excel_row=item["source_row_number"],
            row_key=item["row_key"],
            rmu=item["rmu"],
            values=tuple(values),
            analysis_detail=item["analysis_detail"],
            suggested_comment=item.get("suggested_comment", ""),
            legacy_row_keys=tuple(item.get("legacy_row_keys") or ()),
        ))

    inputs = (ioa_path.resolve(), adms_sld_path.resolve(), standard_workbook_path.resolve())
    return DBSmartReport(
        source_path=ioa_path.resolve(),
        sheet_name="Calculated · STANDARD → ADMS validation + ZENON implementation hint",
        source_hash=_combined_hash(inputs),
        group_definitions=SIGNAL_MAPPING_GROUPS,
        columns=SIGNAL_MAPPING_COLUMNS,
        rows=tuple(rows),
        analysis_column=ANALYSIS_COLUMN,
        input_paths=inputs,
    )

def build_signal_mapping_report_from_store(store) -> DBSmartReport:
    ioa = store.source_path("ioa")
    adms_sld = store.source_path("adms_sld")
    standard = standard_reference_path()
    missing = []
    if not ioa:
        missing.append("ZENON-ADMS-IOA.csv")
    if not adms_sld:
        missing.append("ADMS-SLD.csv")
    if not standard.exists():
        missing.append("active IOA STANDARD.xlsx")
    if missing:
        raise FileNotFoundError("Signal Mapping Review requires: " + ", ".join(missing))
    return build_signal_mapping_report(
        ioa, adms_sld, standard,
        ioa_overrides=store.source_column_overrides("ioa"),
        adms_sld_overrides=store.source_column_overrides("adms_sld"),
        standard_overrides=store.source_column_overrides("standard_reference"),
        ioa_sheet_name=store.source_sheet_name("ioa") if hasattr(store, "source_sheet_name") else None,
        adms_sld_sheet_name=store.source_sheet_name("adms_sld") if hasattr(store, "source_sheet_name") else None,
    )


# Compatibility aliases for older plugins/tests.  Legacy DB-smart worksheet
# discovery is no longer used by the application.
def find_db_smart_sheet(sheet_names: Iterable[str]) -> str | None:
    return None


def read_db_smart_report(path: Path) -> DBSmartReport:  # pragma: no cover - explicit migration guard
    raise RuntimeError(
        "Signal Mapping Review is calculated from ZENON-ADMS-IOA.csv + ADMS-SLD.csv + the active application IOA STANDARD.xlsx; "
        "the legacy DB-smart worksheet is no longer a source."
    )
