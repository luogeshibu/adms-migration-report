"""Tabular file parsing and feeder-normalization helpers. No UI dependencies."""
from __future__ import annotations

import csv
import re
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook


def normalize_key(value) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if text.endswith(".0") and text[:-2].isdigit():
        text = text[:-2]
    return text


def clean(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def local_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1] if "}" in tag else tag


def direct_child_text(elem: ET.Element, child_name: str) -> str:
    for child in elem:
        if local_name(child.tag) == child_name:
            return (child.text or "").strip()
    return ""


CABINET_TYPE_RE = re.compile(r"(?<![A-Za-z0-9])(\d+L(?:\d+T)?)(?![A-Za-z0-9])", re.I)
FEEDER_TOKEN_RE = re.compile(r"[A-Z0-9]+", re.I)


def parse_destination(value: str) -> tuple[str, str, str]:
    """Parse SubstituteDestination using the proven RMU extraction rule.

    Any zenOn project/driver prefix before the final ``#`` is ignored. The last
    hyphen-delimited token is the RMU and everything before it is the feeder.
    """
    value = clean(value)
    if not value:
        return "", "", ""
    normalized = value.rsplit("#", 1)[-1].strip()
    if "-" not in normalized:
        return normalized, "", ""
    feeder, rmu = normalized.rsplit("-", 1)
    return normalized, feeder.strip(), rmu.strip()


def _feeder_tokens(value: str) -> list[str]:
    return [x.upper() for x in FEEDER_TOKEN_RE.findall(clean(value))]


def normalize_feeder(value: str) -> str:
    """Normalize feeder text without losing token boundaries."""
    return "-".join(_feeder_tokens(value))


def feeder_matches_site(feeder: str, site_name: str | None, allowed_feeders: Iterable[str] | None = None) -> bool:
    """Return whether a feeder belongs to the selected site.

    Matching is deliberately token-aware: site ``ABN`` does **not** match token
    ``ABN2``. This is important for combined files such as ABN-ABN2.XML.

    SE feeder values are an additional positive signal. A short SE feeder like
    ``ABN2-16`` matches ``JED-NTH-ABN2-16`` by suffix, which also supports sites
    whose repository folder name is not identical to the feeder token.
    """
    if not site_name:
        return True
    feeder_norm = normalize_feeder(feeder)
    if not feeder_norm:
        return False

    feeder_tokens = feeder_norm.split("-")
    site_tokens = _feeder_tokens(site_name)
    if site_tokens:
        # Repository folders may have an ordering prefix such as ``1-ABH``.
        # Try the full label and each suffix while keeping exact token
        # boundaries, so ABN still never matches ABN2.
        site_candidates = [site_tokens[i:] for i in range(len(site_tokens)) if site_tokens[i:]]
        site_candidates.sort(key=len, reverse=True)
        for candidate in site_candidates:
            n = len(candidate)
            for i in range(0, len(feeder_tokens) - n + 1):
                if feeder_tokens[i:i+n] == candidate:
                    return True

    for candidate in allowed_feeders or ():
        candidate_norm = normalize_feeder(candidate)
        if not candidate_norm:
            continue
        if feeder_norm == candidate_norm or feeder_norm.endswith("-" + candidate_norm):
            return True
    return False



def read_csv_rows(path: Path) -> list[dict]:
    raw = path.read_bytes()
    text = None
    for encoding in ("utf-8-sig", "utf-16", "gb18030", "latin1"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            pass
    if text is None:
        raise ValueError(f"Unable to decode {path.name}")
    return list(csv.DictReader(text.splitlines()))


def read_excel_rows(path: Path) -> list[dict]:
    """Read the active worksheet and release the workbook deterministically.

    ``read_only=True`` keeps the XLSX zip archive open while rows are streamed.
    On Windows that handle prevents ``TemporaryDirectory`` / project cleanup
    from deleting an imported workbook until ``Workbook.close()`` is called.
    Always close in ``finally`` so normal returns, empty workbooks and parser
    exceptions all release the file immediately.
    """
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.active
        rows = ws.iter_rows(values_only=True)
        try:
            headers = [clean(x) for x in next(rows)]
        except StopIteration:
            return []
        result = []
        for values in rows:
            if not any(v is not None and clean(v) for v in values):
                continue
            result.append({headers[i]: values[i] if i < len(values) else None for i in range(len(headers))})
        return result
    finally:
        wb.close()
