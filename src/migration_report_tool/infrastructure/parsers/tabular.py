"""Low-level CSV/XLSX readers plus canonical schema enforcement."""
from __future__ import annotations

import csv
from dataclasses import dataclass
from pathlib import Path
from typing import Mapping

from openpyxl import load_workbook

from ...config.sources import schema_for
from ...domain.schema import (
    SchemaValidationError,
    SchemaValidationResult,
    canonicalize_rows,
    resolve_schema,
)


@dataclass(frozen=True)
class MappedRows:
    rows: tuple[dict, ...]
    validation: SchemaValidationResult


def _decode_csv(path: Path) -> str:
    raw = Path(path).read_bytes()
    for encoding in ("utf-8-sig", "utf-16", "gb18030", "latin1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError(f"Unable to decode {Path(path).name}")


def read_csv_raw(path: Path) -> tuple[list[str], list[dict]]:
    reader = csv.DictReader(_decode_csv(Path(path)).splitlines())
    headers = [str(x or "").strip() for x in (reader.fieldnames or [])]
    return headers, list(reader)


@dataclass(frozen=True)
class ExcelSheetInfo:
    name: str
    rows: int
    columns: int
    usable: bool


def _clean_excel_row(values) -> list[str]:
    return [str(x or "").replace("\ufeff", "").strip() for x in values]


def _sheet_is_usable(ws, scan_rows: int = 12) -> bool:
    """Return True when a sheet looks like a normal header + data table.

    AUTO intentionally stays conservative: workbook order is respected and the
    first usable sheet wins.  If none qualifies, the workbook's first sheet is
    still used so the reviewer can remap its physical fields manually.
    """
    nonempty_rows = 0
    first_nonempty = None
    for values in ws.iter_rows(min_row=1, max_row=min(max(1, ws.max_row), scan_rows), values_only=True):
        cleaned = _clean_excel_row(values)
        if not any(cleaned):
            continue
        nonempty_rows += 1
        if first_nonempty is None:
            first_nonempty = cleaned
        if nonempty_rows >= 2:
            break
    return bool(first_nonempty and any(first_nonempty) and nonempty_rows >= 2)


def list_excel_sheets(path: Path) -> tuple[ExcelSheetInfo, ...]:
    """Return workbook sheet metadata in physical workbook order."""
    wb = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        result = []
        for name in wb.sheetnames:
            ws = wb[name]
            result.append(ExcelSheetInfo(name, int(ws.max_row or 0), int(ws.max_column or 0), _sheet_is_usable(ws)))
        return tuple(result)
    finally:
        wb.close()


def resolve_excel_sheet_name(path: Path, preferred_sheet: str | None = None) -> str:
    """Resolve a site-level sheet selection.

    A saved manual sheet wins when it still exists. AUTO otherwise selects the
    first usable sheet in workbook order, falling back to the first physical
    sheet for intentionally unusual workbooks that will be mapped manually.
    """
    wb = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        names = list(wb.sheetnames)
        if not names:
            raise ValueError(f"Workbook contains no sheets: {Path(path).name}")
        wanted = str(preferred_sheet or "").strip()
        if wanted:
            exact = next((name for name in names if name.casefold() == wanted.casefold()), None)
            if exact:
                return exact
        for name in names:
            if _sheet_is_usable(wb[name]):
                return name
        return names[0]
    finally:
        wb.close()


def resolve_source_excel_sheet_name(source_type: str, path: Path, preferred_sheet: str | None = None) -> str:
    """Resolve the effective Excel sheet for one source role.

    A site-level/manual selection always wins.  When AUTO is active, a source
    schema may declare a preferred business sheet (for example ADMS SLD uses
    ``ADMS-SLD主设备``).  If that sheet is absent, the historical conservative
    first-usable-sheet behavior remains the fallback.
    """
    path = Path(path)
    preferred = str(preferred_sheet or "").strip()
    if preferred:
        return resolve_excel_sheet_name(path, preferred)
    schema = schema_for(source_type)
    schema_sheet = str(schema.sheet_name or "").strip() if schema else ""
    return resolve_excel_sheet_name(path, schema_sheet or None)


def read_excel_raw(path: Path, *, sheet_name: str | None = None) -> tuple[list[str], list[dict]]:
    wb = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        names = list(wb.sheetnames)
        if not names:
            return [], []
        wanted = str(sheet_name or "").strip()
        selected = next((name for name in names if wanted and name.casefold() == wanted.casefold()), None)
        if selected is None:
            selected = next((name for name in names if _sheet_is_usable(wb[name])), names[0])
        ws = wb[selected]
        iterator = ws.iter_rows(values_only=True)
        try:
            headers = _clean_excel_row(next(iterator))
        except StopIteration:
            return [], []
        rows: list[dict] = []
        for values in iterator:
            if not any(v is not None and str(v).strip() for v in values):
                continue
            rows.append({headers[i]: values[i] if i < len(values) else None for i in range(len(headers))})
        return headers, rows
    finally:
        wb.close()


def read_standard_raw(path: Path, *, sheet_name: str = "STANDARD", scan_rows: int = 25, overrides: Mapping[str, str] | None = None) -> tuple[list[str], list[dict]]:
    """Read STANDARD by header names, not fixed coordinates."""
    wb = load_workbook(path, read_only=False, data_only=True, keep_links=False)
    try:
        matches = [name for name in wb.sheetnames if name.strip().casefold() == sheet_name.casefold()]
        if not matches:
            raise ValueError(f"{sheet_name} sheet not found in {Path(path).name}.")
        ws = wb[matches[0]]
        schema = schema_for("standard_reference")
        best = None
        for row_index in range(1, min(ws.max_row, scan_rows) + 1):
            headers = [str(ws.cell(row_index, c).value or "").replace("\ufeff", "").strip() for c in range(1, ws.max_column + 1)]
            validation = resolve_schema(schema, headers, overrides)
            if not validation.errors:
                best = (row_index, headers, validation)
                break
        if not best:
            # Give the user the most useful error using the first non-empty header row.
            for row_index in range(1, min(ws.max_row, scan_rows) + 1):
                headers = [str(ws.cell(row_index, c).value or "").replace("\ufeff", "").strip() for c in range(1, ws.max_column + 1)]
                if any(headers):
                    validation = resolve_schema(schema, headers, overrides)
                    raise SchemaValidationError(validation, f"{Path(path).name} / {sheet_name}")
            raise ValueError(f"No header row found in {Path(path).name} / {sheet_name}.")
        row_index, headers, _validation = best
        rows: list[dict] = []
        for r in range(row_index + 1, ws.max_row + 1):
            values = [ws.cell(r, c).value for c in range(1, len(headers) + 1)]
            if not any(v is not None and str(v).strip() for v in values):
                continue
            rows.append({headers[i]: values[i] if i < len(values) else None for i in range(len(headers))})
        return headers, rows
    finally:
        wb.close()


def validate_source_file(source_type: str, path: Path, overrides: Mapping[str, str] | None = None, *, sheet_name: str | None = None) -> SchemaValidationResult | None:
    schema = schema_for(source_type)
    if schema is None:
        return None
    path = Path(path)
    if source_type == "standard_reference":
        headers, _ = read_standard_raw(path, sheet_name=schema.sheet_name or "STANDARD", scan_rows=schema.header_scan_rows, overrides=overrides)
    elif path.suffix.lower() == ".csv":
        headers, _ = read_csv_raw(path)
    elif path.suffix.lower() in {".xlsx", ".xlsm"}:
        effective_sheet = resolve_source_excel_sheet_name(source_type, path, sheet_name)
        headers, _ = read_excel_raw(path, sheet_name=effective_sheet)
    else:
        return None
    return resolve_schema(schema, headers, overrides)


def read_mapped_rows(
    source_type: str,
    path: Path,
    overrides: Mapping[str, str] | None = None,
    *,
    strict: bool = True,
    sheet_name: str | None = None,
) -> MappedRows:
    schema = schema_for(source_type)
    if schema is None:
        raise ValueError(f"No tabular schema is registered for source type: {source_type}")
    path = Path(path)
    if source_type == "standard_reference":
        headers, rows = read_standard_raw(path, sheet_name=schema.sheet_name or "STANDARD", scan_rows=schema.header_scan_rows, overrides=overrides)
    elif path.suffix.lower() == ".csv":
        headers, rows = read_csv_raw(path)
    elif path.suffix.lower() in {".xlsx", ".xlsm"}:
        effective_sheet = resolve_source_excel_sheet_name(source_type, path, sheet_name)
        headers, rows = read_excel_raw(path, sheet_name=effective_sheet)
    else:
        raise ValueError(f"Unsupported tabular source format: {path.name}")
    validation = resolve_schema(schema, headers, overrides)
    if strict and validation.errors:
        raise SchemaValidationError(validation, path.name)
    return MappedRows(tuple(canonicalize_rows(rows, validation)), validation)
