"""Site-local, fully configurable Excel/CSV comparison engine.

v0.8.178 replaces the *fixed five-source input contract* of Equipment Data
Review with a site-local configuration.  The downstream human review workflow
(status, comments, resolutions, lifecycle, audit and reporting) continues to
consume the same row-shaped payload.

Important design rules:
* Any number of CSV/XLSX/XLSM source tables may participate.
* Physical filenames and worksheet names are not business contracts.
* Every source chooses its own key/index column.
* Comparison fields are logical rules configured by the reviewer; each rule may
  bind any physical column from any source.
* All physical source columns are visible by default.  Hidden columns are saved
  per site and per source.
* Comparison semantics are configurable per site and per rule. A site chooses
  a default mode, and each Comparison Rule may inherit that default or override
  it with strict blank-aware equality or legacy blank-ignoring equality.
* Source titles default to the filename stem and are reviewer-editable.
"""
from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
import csv
import hashlib
import json
import re
import uuid

from openpyxl import load_workbook

from ..analysis import compare_consistency, compare_strict_consistency
from ..parsers import clean
from ..infrastructure.database.global_settings_store import (
    equipment_comparison_profiles as _global_equipment_profiles,
    equipment_comparison_profile as _global_equipment_profile,
    replace_equipment_comparison_profile as _replace_global_equipment_profile,
    delete_equipment_comparison_profile as _delete_global_equipment_profile,
)

CONFIG_KEY = "equipment_comparison_config_v1"
CONFIG_VERSION = 4
PROFILE_LINK_KEY = "equipment_comparison_profile_link_v1"
PROFILE_PAYLOAD_VERSION = 2
LIVE_METADATA_KEY = "configurable_live_source_metadata_v1"

COMPARISON_MODE_DEFAULT = "default"
COMPARISON_MODE_STRICT = "strict"
COMPARISON_MODE_IGNORE_BLANK = "ignore_blank"
VALID_SITE_COMPARISON_MODES = {COMPARISON_MODE_STRICT, COMPARISON_MODE_IGNORE_BLANK}
VALID_RULE_COMPARISON_MODES = {COMPARISON_MODE_DEFAULT, *VALID_SITE_COMPARISON_MODES}


def _site_comparison_mode(value: object) -> str:
    mode = clean(value)
    return mode if mode in VALID_SITE_COMPARISON_MODES else COMPARISON_MODE_STRICT


def _rule_comparison_mode(rule: dict, site_default: str) -> str:
    mode = clean((rule or {}).get("comparison_mode"))
    if mode in VALID_SITE_COMPARISON_MODES:
        return mode
    return _site_comparison_mode(site_default)


@dataclass(frozen=True)
class PhysicalColumn:
    id: str
    header: str
    label: str
    index: int


@dataclass(frozen=True)
class TableStructure:
    path: Path
    sheet_name: str
    header_row: int
    columns: tuple[PhysicalColumn, ...]


@dataclass(frozen=True)
class LoadedTable:
    structure: TableStructure
    rows: tuple[dict[str, object], ...]


def _decode_csv(path: Path) -> str:
    raw = Path(path).read_bytes()
    for encoding in ("utf-8-sig", "utf-16", "gb18030", "latin1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError(f"Unable to decode {Path(path).name}")


def _field_id(header: str, occurrence: int) -> str:
    # Header + occurrence stays stable when unrelated columns move around.
    token = (str(header or "").strip().casefold() or "<blank>") + f"#{int(occurrence)}"
    return "c_" + hashlib.sha1(token.encode("utf-8")).hexdigest()[:12]


def _make_columns(headers: list[object]) -> tuple[PhysicalColumn, ...]:
    total: dict[str, int] = {}
    for raw in headers:
        key = clean(raw).casefold()
        total[key] = total.get(key, 0) + 1
    seen: dict[str, int] = {}
    result: list[PhysicalColumn] = []
    for index, raw in enumerate(headers):
        header = clean(raw) or f"Column {index + 1}"
        key = clean(raw).casefold()
        occurrence = seen.get(key, 0) + 1
        seen[key] = occurrence
        label = header if total.get(key, 0) <= 1 else f"{header} [{occurrence}]"
        result.append(PhysicalColumn(_field_id(header, occurrence), header, label, index))
    return tuple(result)


def list_sheets(path: Path) -> tuple[str, ...]:
    path = Path(path)
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        return ()
    wb = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        return tuple(wb.sheetnames)
    finally:
        wb.close()


def _first_usable_sheet(path: Path) -> str:
    wb = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        if not wb.sheetnames:
            raise ValueError(f"Workbook contains no sheets: {path.name}")
        for name in wb.sheetnames:
            ws = wb[name]
            if int(ws.max_row or 0) >= 2 and int(ws.max_column or 0) >= 1:
                return name
        return wb.sheetnames[0]
    finally:
        wb.close()


def inspect_table(path: Path, *, sheet_name: str = "", header_row: int = 1) -> TableStructure:
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(path)
    header_row = max(1, int(header_row or 1))
    suffix = path.suffix.lower()
    if suffix == ".csv":
        reader = csv.reader(_decode_csv(path).splitlines())
        rows = list(reader)
        if len(rows) < header_row:
            raise ValueError(f"Header row {header_row} does not exist in {path.name}")
        headers = rows[header_row - 1]
        return TableStructure(path, "", header_row, _make_columns(headers))
    if suffix not in {".xlsx", ".xlsm"}:
        raise ValueError("Only CSV, XLSX and XLSM are supported")
    selected = clean(sheet_name) or _first_usable_sheet(path)
    wb = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        if selected not in wb.sheetnames:
            raise ValueError(f"Worksheet not found: {selected}")
        ws = wb[selected]
        values = next(ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True), ())
        return TableStructure(path, selected, header_row, _make_columns(list(values)))
    finally:
        wb.close()


def load_table(path: Path, *, sheet_name: str = "", header_row: int = 1) -> LoadedTable:
    structure = inspect_table(path, sheet_name=sheet_name, header_row=header_row)
    columns = structure.columns
    suffix = structure.path.suffix.lower()
    output: list[dict[str, object]] = []
    if suffix == ".csv":
        reader = csv.reader(_decode_csv(structure.path).splitlines())
        for row_index, values in enumerate(reader, start=1):
            if row_index <= structure.header_row:
                continue
            if not any(clean(value) for value in values):
                continue
            output.append({column.id: values[column.index] if column.index < len(values) else "" for column in columns})
        return LoadedTable(structure, tuple(output))

    wb = load_workbook(structure.path, read_only=True, data_only=True, keep_links=False)
    try:
        ws = wb[structure.sheet_name]
        for values in ws.iter_rows(min_row=structure.header_row + 1, values_only=True):
            if not any(clean(value) for value in values):
                continue
            output.append({column.id: values[column.index] if column.index < len(values) else "" for column in columns})
    finally:
        wb.close()
    return LoadedTable(structure, tuple(output))


def _source_id() -> str:
    return "src_" + uuid.uuid4().hex[:12]


def _rule_id() -> str:
    return "cmp_" + uuid.uuid4().hex[:12]


def empty_config() -> dict:
    return {
        "version": CONFIG_VERSION,
        # Backward compatibility: v0.8.191 made configurable comparison strict.
        # Existing sites therefore remain strict until the reviewer explicitly
        # chooses the legacy Ignore Blank mode.
        "default_comparison_mode": COMPARISON_MODE_STRICT,
        "sources": [],
        "comparisons": [],
        "updated_at": "",
    }


def normalize_config(config: dict | None) -> dict:
    raw = dict(config or {})
    output = empty_config()
    output["version"] = CONFIG_VERSION
    output["default_comparison_mode"] = _site_comparison_mode(raw.get("default_comparison_mode"))
    seen_sources = set()
    sources = []
    for item in raw.get("sources", []) or []:
        source = dict(item or {})
        source_id = clean(source.get("id")) or _source_id()
        if source_id in seen_sources:
            source_id = _source_id()
        seen_sources.add(source_id)
        sources.append({
            "id": source_id,
            "title": clean(source.get("title")),
            "path": str(source.get("path") or "").strip(),
            "path_mode": clean(source.get("path_mode")) or "absolute",
            "sheet_name": clean(source.get("sheet_name")),
            "header_row": max(1, int(source.get("header_row") or 1)),
            "key_column": clean(source.get("key_column")),
            "hidden_columns": sorted({clean(x) for x in (source.get("hidden_columns") or []) if clean(x)}),
            "enabled": bool(source.get("enabled", True)),
            "selection_mode": clean(source.get("selection_mode")) if clean(source.get("selection_mode")) in {"pinned", "latest_family"} else "pinned",
            "family_key": clean(source.get("family_key")),
            "family_suffix": clean(source.get("family_suffix")),
        })
    output["sources"] = sources

    source_ids = {item["id"] for item in sources}
    seen_rules = set()
    comparisons = []
    for item in raw.get("comparisons", []) or []:
        rule = dict(item or {})
        rule_id = clean(rule.get("id")) or _rule_id()
        if rule_id in seen_rules:
            rule_id = _rule_id()
        seen_rules.add(rule_id)
        title = clean(rule.get("title")) or "Comparison"
        bindings = {
            clean(source_id): clean(column_id)
            for source_id, column_id in dict(rule.get("bindings") or {}).items()
            if clean(source_id) in source_ids and clean(column_id)
        }
        mode = clean(rule.get("comparison_mode"))
        if mode not in VALID_RULE_COMPARISON_MODES:
            mode = COMPARISON_MODE_DEFAULT
        comparisons.append({"id": rule_id, "title": title, "comparison_mode": mode, "bindings": bindings})
    output["comparisons"] = comparisons
    output["updated_at"] = clean(raw.get("updated_at"))
    return output


def get_config(store, *, bootstrap: bool = False) -> dict:
    config = normalize_config((store.config or {}).get(CONFIG_KEY) or {})
    if bootstrap and not config["sources"]:
        config = bootstrap_legacy_sources(store)
    return config


def save_config(store, config: dict, modified_by: str = "") -> dict:
    normalized = normalize_config(config)

    # v0.8.182: configurable sources are references to live files, not archived
    # source snapshots.  Promote any legacy internal project/source_files path to
    # the reviewer-facing live/original path before persisting the configuration.
    # Existing historical snapshot files are left untouched for audit/backward
    # compatibility; no new copy is created by this workflow.
    for source in normalized.get("sources", []):
        stored = _stored_source_path(store, source)
        if stored is None or not _path_is_inside_project_data(store, stored):
            continue
        effective = source_path(store, source)
        if effective is None or not effective.exists() or _path_is_inside_project_data(store, effective):
            continue
        encoded, path_mode = encode_source_path(store, effective)
        source["path"] = encoded
        source["path_mode"] = path_mode
        source["family_key"] = file_family_key(effective)
        source["family_suffix"] = effective.suffix.lower()

    normalized["updated_at"] = datetime.now().isoformat(timespec="seconds")
    old = json.dumps(normalize_config((store.config or {}).get(CONFIG_KEY) or {}), ensure_ascii=False, sort_keys=True)
    new = json.dumps(normalized, ensure_ascii=False, sort_keys=True)
    store.config[CONFIG_KEY] = normalized
    store.save_config()
    if old != new and hasattr(store, "db"):
        now = datetime.now().isoformat(timespec="seconds")
        try:
            store.db.execute(
                "INSERT INTO changes(rmu,field_name,old_value,new_value,reason,modified_by,modified_at) VALUES(?,?,?,?,?,?,?)",
                ("EQUIPMENT_COMPARISON", "comparison.configuration", old, new,
                 "Site-local configurable equipment comparison", modified_by or "system", now),
            )
            store.db.commit()
        except Exception:
            pass
    return normalized


def _profile_safe_config(config: dict | None) -> dict:
    """Return reusable logical comparison configuration with no physical paths.

    A profile intentionally preserves source roles, file-family hints, worksheet /
    header defaults, Key/Index selection, visibility defaults and comparison
    bindings.  It never stores an absolute or site-relative source path, so
    applying a profile at another station cannot read the originating station.
    """
    normalized = normalize_config(config or {})
    for source in normalized.get("sources", []):
        source["path"] = ""
        source["path_mode"] = "absolute"
    normalized["updated_at"] = ""
    return normalized


def list_comparison_profiles() -> list[dict]:
    return [
        {
            "name": clean(item.get("name")),
            "modified_by": clean(item.get("modified_by")),
            "modified_at": clean(item.get("modified_at")),
        }
        for item in _global_equipment_profiles()
        if clean(item.get("name"))
    ]


def get_comparison_profile(profile_name: str) -> dict | None:
    item = _global_equipment_profile(clean(profile_name))
    if not item:
        return None
    return {
        "name": clean(item.get("name")),
        "config": _profile_safe_config((item.get("payload") or {}).get("equipment_config") or {}),
        "modified_by": clean(item.get("modified_by")),
        "modified_at": clean(item.get("modified_at")),
    }


def save_comparison_profile(profile_name: str, config: dict, modified_by: str = "") -> dict:
    name = clean(profile_name)
    if not name:
        raise ValueError("Profile name is required")
    payload = {
        "version": PROFILE_PAYLOAD_VERSION,
        "equipment_config": _profile_safe_config(config),
    }
    saved = _replace_global_equipment_profile(name, payload, modified_by or "system")
    return {
        "name": clean(saved.get("name")),
        "config": _profile_safe_config((saved.get("payload") or {}).get("equipment_config") or {}),
        "modified_by": clean(saved.get("modified_by")),
        "modified_at": clean(saved.get("modified_at")),
    }


def delete_comparison_profile(profile_name: str, modified_by: str = "") -> bool:
    return bool(_delete_global_equipment_profile(clean(profile_name), modified_by or "system"))


def get_profile_link(store) -> dict:
    raw = dict((store.config or {}).get(PROFILE_LINK_KEY) or {})
    mode = clean(raw.get("mode"))
    if mode not in {"local", "inherit"}:
        mode = "local"
    return {
        "mode": mode,
        "profile_name": clean(raw.get("profile_name")),
        "profile_modified_at": clean(raw.get("profile_modified_at")),
        "synced_at": clean(raw.get("synced_at")),
    }


def save_profile_link(
    store, *, mode: str = "local", profile_name: str = "",
    profile_modified_at: str = "", synced_at: str = ""
) -> dict:
    mode = "inherit" if clean(mode) == "inherit" and clean(profile_name) else "local"
    if mode == "local":
        store.config.pop(PROFILE_LINK_KEY, None)
        store.save_config()
        return {"mode": "local", "profile_name": "", "profile_modified_at": "", "synced_at": ""}
    value = {
        "mode": "inherit",
        "profile_name": clean(profile_name),
        "profile_modified_at": clean(profile_modified_at),
        "synced_at": clean(synced_at) or datetime.now().isoformat(timespec="seconds"),
    }
    store.config[PROFILE_LINK_KEY] = value
    store.save_config()
    return value


SUPPORTED_EXTENSIONS = {".csv", ".xlsx", ".xlsm"}


def site_root(store) -> Path | None:
    text = str((store.config or {}).get("repository_path") or "").strip()
    if not text:
        return None
    root = Path(text)
    return root if root.exists() and root.is_dir() else None


def file_family_key(path_or_name: Path | str) -> str:
    """Return a forgiving logical family key for user-managed table versions.

    Generic comparison sources intentionally do not require the historical strict
    ``-V1`` naming contract.  Common copy/version/date suffixes are removed so
    names such as ``NEWSYSTEM-SLD.xlsx``, ``NEWSYSTEM-SLD (2).xlsx`` and
    ``NEWSYSTEM-SLD-V3.xlsx`` can be treated as one selectable family.
    """
    stem = Path(str(path_or_name)).stem.strip()
    previous = None
    while previous != stem:
        previous = stem
        stem = re.sub(r"\s*\(\d+\)$", "", stem, flags=re.IGNORECASE).strip()
        stem = re.sub(r"(?:[-_.\s]+)(?:v|ver|version|rev)[-_.\s]*\d+(?:\.\d+)*$", "", stem, flags=re.IGNORECASE).strip()
        stem = re.sub(r"(?:[-_.\s]+)20\d{2}[-_.]?\d{2}[-_.]?\d{2}(?:[-_.]?\d{4,6})?$", "", stem, flags=re.IGNORECASE).strip()
        stem = re.sub(r"(?:[-_.\s]+)(?:copy|final|latest)$", "", stem, flags=re.IGNORECASE).strip()
    token = re.sub(r"[^a-z0-9]+", "", stem.casefold())
    return token or re.sub(r"[^a-z0-9]+", "", Path(str(path_or_name)).stem.casefold())


def _generic_version_tuple(path: Path) -> tuple[int, ...]:
    stem = path.stem.strip()
    match = re.search(r"(?:^|[-_.\s])(?:v|ver|version|rev)[-_.\s]*(\d+(?:\.\d+)*)$", stem, flags=re.IGNORECASE)
    if not match:
        return ()
    try:
        return tuple(int(part) for part in match.group(1).split("."))
    except Exception:
        return ()


def scan_site_tabular_files(store, *, max_depth: int | None = None) -> tuple[Path, ...]:
    """Discover supported source files under the active site folder.

    Root-only sites continue to work, while optional ``Equipment`` /
    ``SignalMapping`` subfolders are also discovered.  Discovery is inventory
    only: a newly seen file is never silently added to a comparison.
    """
    root = site_root(store)
    if root is None:
        return ()
    found: list[Path] = []
    try:
        root_parts = len(root.resolve().parts)
    except Exception:
        root_parts = len(root.parts)
    for candidate in root.rglob("*"):
        try:
            if not candidate.is_file() or candidate.suffix.lower() not in SUPPORTED_EXTENSIONS:
                continue
            if candidate.name.startswith("~$"):
                continue
            rel = candidate.relative_to(root)
            if max_depth is not None and len(rel.parts) > max_depth + 1:
                continue
            if any(part.startswith(".") or part.lower() in {"archive", "backup", "backups", "__pycache__"} for part in rel.parts[:-1]):
                continue
            found.append(candidate)
        except (OSError, ValueError):
            continue
    def key(path: Path):
        try:
            mtime = path.stat().st_mtime_ns
        except OSError:
            mtime = 0
        return (file_family_key(path), path.name.casefold(), -mtime, str(path).casefold())
    return tuple(sorted(found, key=key))


def family_candidates(store, family_key: str, *, suffix: str = "") -> tuple[Path, ...]:
    family_key = clean(family_key).casefold()
    candidates = [
        path for path in scan_site_tabular_files(store)
        if file_family_key(path) == family_key and (not suffix or path.suffix.lower() == suffix.lower())
    ]
    def rank(path: Path):
        version = _generic_version_tuple(path)
        padded = version + (0,) * max(0, 8 - len(version))
        try:
            mtime = path.stat().st_mtime_ns
        except OSError:
            mtime = 0
        return (1 if version else 0, padded, mtime, path.name.casefold())
    candidates.sort(key=rank, reverse=True)
    return tuple(candidates)


def _stored_source_path(store, source: dict) -> Path | None:
    text = str((source or {}).get("path") or "").strip()
    if not text:
        return None
    mode = clean((source or {}).get("path_mode")).lower()
    if mode == "site_relative":
        base = site_root(store)
        if base:
            return base / text
    return Path(text)


def _legacy_role_for_source(source: dict, stored: Path | None = None) -> str:
    """Recover the old logical role only to escape a historical workspace copy.

    The configurable engine does not otherwise depend on these names.  Filename
    prefix fallback matters when the reviewer already renamed a bootstrapped
    module title before upgrading to direct-read mode.
    """
    title = clean((source or {}).get("title")).casefold()
    by_title = {
        "se": "se_list",
        "se equipment": "se_list",
        "zenon db": "zenon_db",
        "zenon sld": "zenon_sld",
        "adms db": "adms_db",
        "adms sld": "adms_sld",
    }.get(title, "")
    if by_title:
        return by_title
    name = (stored.name if stored is not None else Path(str((source or {}).get("path") or "")).name).casefold()
    for prefix, role in (
        ("se_list_", "se_list"), ("zenon_db_", "zenon_db"),
        ("zenon_sld_", "zenon_sld"), ("adms_db_", "adms_db"),
        ("adms_sld_", "adms_sld"),
    ):
        if name.startswith(prefix):
            return role
    return ""


def source_path(store, source: dict) -> Path | None:
    """Resolve the live physical file for one configured source.

    v0.8.182 is *direct-read only* for configurable sources: the returned path
    points at the reviewer's real CSV/XLSX/XLSM file.  The function never copies
    a workbook into Project Data. ``pinned`` reads that exact live file in place;
    ``latest_family`` follows the newest sibling in the configured file family.

    For projects migrated from the historical five-source engine, a stored path
    may still refer to ``project/source_files/<timestamped copy>``.  When that is
    detected, the resolver transparently rebinds to the remembered/original live
    repository file whenever possible, so the UI and comparison engine no longer
    read stale workspace snapshots.
    """
    stored = _stored_source_path(store, source)
    mode = clean((source or {}).get("selection_mode")).lower() or "pinned"

    # Existing v0.8.178-v0.8.181 configs can contain a timestamped legacy copy.
    # Prefer the true live/original file without mutating/deleting history.
    if stored is not None and _path_is_inside_project_data(store, stored):
        role = _legacy_role_for_source(source, stored)
        live = _legacy_path(store, role) if role else None
        if live is not None and live.exists() and not _path_is_inside_project_data(store, live):
            stored = live

    if mode != "latest_family":
        return stored
    family = clean((source or {}).get("family_key")) or (file_family_key(stored) if stored else "")
    suffix = clean((source or {}).get("family_suffix")) or (stored.suffix.lower() if stored else "")
    candidates = family_candidates(store, family, suffix=suffix) if family else ()
    return candidates[0] if candidates else stored



SIGNAL_ASSIGNMENTS_KEY = "signal_mapping_source_assignments_v1"
SIGNAL_ASSIGNMENT_ROLES = ("ioa", "adms_sld")


def get_signal_source_assignments(store) -> dict[str, dict]:
    raw = dict((store.config or {}).get(SIGNAL_ASSIGNMENTS_KEY) or {})
    result: dict[str, dict] = {}
    for role in SIGNAL_ASSIGNMENT_ROLES:
        item = dict(raw.get(role) or {})
        if not item:
            continue
        selection_mode = clean(item.get("selection_mode"))
        result[role] = {
            "path": str(item.get("path") or "").strip(),
            "path_mode": clean(item.get("path_mode")) or "absolute",
            "selection_mode": selection_mode if selection_mode in {"pinned", "latest_family"} else "pinned",
            "family_key": clean(item.get("family_key")),
            "family_suffix": clean(item.get("family_suffix")),
            "title": clean(item.get("title")),
        }
    return result


def save_signal_source_assignment(store, role: str, path: Path | str | None, *, selection_mode: str = "pinned", title: str = "") -> None:
    role = clean(role)
    if role not in SIGNAL_ASSIGNMENT_ROLES:
        raise ValueError(f"Unsupported Signal Mapping source role: {role}")
    root = dict((store.config or {}).get(SIGNAL_ASSIGNMENTS_KEY) or {})
    if path is None:
        root.pop(role, None)
    else:
        physical = Path(path)
        encoded, path_mode = encode_source_path(store, physical)
        root[role] = {
            "path": encoded,
            "path_mode": path_mode,
            "selection_mode": selection_mode if selection_mode in {"pinned", "latest_family"} else "pinned",
            "family_key": file_family_key(physical),
            "family_suffix": physical.suffix.lower(),
            "title": clean(title) or physical.stem,
        }
    store.config[SIGNAL_ASSIGNMENTS_KEY] = root
    store.config["validation_required_after_source_import"] = True
    store.save_config()


def resolve_signal_source_assignment(store, role: str) -> Path | None:
    item = get_signal_source_assignments(store).get(clean(role))
    if not item:
        return None
    return source_path(store, item)


def configured_live_source_metadata(store) -> dict[str, dict]:
    """Return stat-only metadata for every currently active direct source.

    This is intentionally cheap: workbook contents are not opened.  It is used
    by the five-second UI watcher to notice in-place edits and AUTO-family
    version changes.  The real file is re-read only after a change is detected.
    """
    result: dict[str, dict] = {}
    config = get_config(store, bootstrap=False)
    for source in config.get("sources", []):
        if not bool(source.get("enabled", True)):
            continue
        source_id = clean(source.get("id"))
        if not source_id:
            continue
        path = source_path(store, source)
        key = f"equipment:{source_id}"
        item = {
            "module": "equipment",
            "source_id": source_id,
            "title": clean(source.get("title")) or (path.stem if path else "Source"),
            "path": str(path.resolve()) if path and path.exists() else str(path or ""),
            "exists": bool(path and path.exists()),
            "size": -1,
            "mtime_ns": -1,
        }
        if path and path.exists():
            try:
                stat = path.stat()
                item["size"] = int(stat.st_size)
                item["mtime_ns"] = int(stat.st_mtime_ns)
            except OSError:
                pass
        result[key] = item

    for role, assignment in get_signal_source_assignments(store).items():
        path = resolve_signal_source_assignment(store, role)
        key = f"signal:{role}"
        item = {
            "module": "signal",
            "source_id": role,
            "title": clean(assignment.get("title")) or role,
            "path": str(path.resolve()) if path and path.exists() else str(path or ""),
            "exists": bool(path and path.exists()),
            "size": -1,
            "mtime_ns": -1,
        }
        if path and path.exists():
            try:
                stat = path.stat()
                item["size"] = int(stat.st_size)
                item["mtime_ns"] = int(stat.st_mtime_ns)
            except OSError:
                pass
        result[key] = item
    return result


def remember_configured_live_source_metadata(store) -> dict[str, dict]:
    metadata = configured_live_source_metadata(store)
    store.config[LIVE_METADATA_KEY] = metadata
    store.save_config()
    return metadata


def configured_live_source_changes(store) -> tuple[tuple[str, ...], tuple]:
    """Return changed direct-source keys plus a deterministic watcher signature.

    An empty baseline means this is the first v0.8.182 observation, not a data
    change.  The baseline is established after the next rebuild/validation.
    """
    current = configured_live_source_metadata(store)
    baseline = dict((store.config or {}).get(LIVE_METADATA_KEY) or {})
    signature = tuple(
        sorted(
            (key, str(item.get("path") or ""), bool(item.get("exists")), int(item.get("size") or -1), int(item.get("mtime_ns") or -1))
            for key, item in current.items()
        )
    )
    if not baseline:
        return (), signature
    changed: list[str] = []
    for key in sorted(set(baseline) | set(current)):
        before = dict(baseline.get(key) or {})
        after = dict(current.get(key) or {})
        before_sig = (str(before.get("path") or ""), bool(before.get("exists")), int(before.get("size") or -1), int(before.get("mtime_ns") or -1))
        after_sig = (str(after.get("path") or ""), bool(after.get("exists")), int(after.get("size") or -1), int(after.get("mtime_ns") or -1))
        if before_sig != after_sig:
            changed.append(key)
    return tuple(changed), signature


def file_pool_usage(store, path: Path) -> tuple[str, ...]:
    """Return human-facing module memberships for one discovered file."""
    try:
        resolved = Path(path).resolve()
    except OSError:
        resolved = Path(path)
    usages: list[str] = []
    config = get_config(store, bootstrap=False)
    for source in config.get("sources", []):
        active = source_path(store, source)
        stored = _stored_source_path(store, source)
        candidates = [p for p in (active, stored) if p is not None]
        if any(_same_path(resolved, p) for p in candidates):
            state = "Equipment" if bool(source.get("enabled", True)) else "Equipment (disabled)"
            if state not in usages:
                usages.append(state)
    for role, label in (("ioa", "Signal · IOA"), ("adms_sld", "Signal · ADMS SLD")):
        active = resolve_signal_source_assignment(store, role)
        if active is not None and _same_path(resolved, active):
            usages.append(label)
    return tuple(usages)


def _same_path(left: Path, right: Path) -> bool:
    try:
        return left.resolve() == Path(right).resolve()
    except OSError:
        return str(left).casefold() == str(right).casefold()

def encode_source_path(store, path: Path) -> tuple[str, str]:
    path = Path(path).resolve()
    base_text = str((store.config or {}).get("repository_path") or "").strip()
    if base_text:
        base = Path(base_text).resolve()
        try:
            rel = path.relative_to(base)
            return str(rel), "site_relative"
        except ValueError:
            pass
    return str(path), "absolute"


def apply_comparison_profile(
    store, profile_name: str, *, current_config: dict | None = None
) -> tuple[dict, dict]:
    """Merge one reusable profile into the current site's physical sources.

    The profile owns the reusable *logical* contract (source roles/order/titles,
    Key/Index defaults, visibility defaults and comparison rules).  Physical file
    paths remain site-local.  Existing local sources are matched by stable id,
    file-family key, then module title; matching local paths are preserved.  For
    a new role the site's file pool is searched by the profile's family key.

    Site-only extra sources are deliberately retained and appended.  This makes
    profile reuse safe for stations that have an extra vendor/system table.
    """
    profile = get_comparison_profile(profile_name)
    if not profile:
        raise ValueError(f"Comparison profile not found: {clean(profile_name) or '<blank>'}")
    template = normalize_config(profile.get("config") or {})
    current = normalize_config(current_config if current_config is not None else get_config(store, bootstrap=False))

    unmatched = list(current.get("sources", []))
    used_current_ids: set[str] = set()
    id_map: dict[str, str] = {}
    merged_sources: list[dict] = []

    def find_match(template_source: dict) -> dict | None:
        template_id = clean(template_source.get("id"))
        family = clean(template_source.get("family_key")).casefold()
        title = clean(template_source.get("title")).casefold()
        for candidate in unmatched:
            if clean(candidate.get("id")) == template_id:
                return candidate
        if family:
            for candidate in unmatched:
                if clean(candidate.get("family_key")).casefold() == family:
                    return candidate
        if title:
            for candidate in unmatched:
                if clean(candidate.get("title")).casefold() == title:
                    return candidate
        return None

    for template_source in template.get("sources", []):
        local = find_match(template_source)
        merged = dict(template_source)
        template_id = clean(template_source.get("id")) or _source_id()
        target_id = clean((local or {}).get("id")) or template_id
        if target_id in {clean(item.get("id")) for item in merged_sources}:
            target_id = _source_id()
        merged["id"] = target_id
        id_map[template_id] = target_id

        # Never inherit another site's path.  Preserve this site's existing file
        # binding when the role already exists; otherwise resolve an AUTO family
        # candidate from this site's own repository/file pool.
        if local is not None:
            merged["path"] = str(local.get("path") or "").strip()
            merged["path_mode"] = clean(local.get("path_mode")) or "absolute"
            # Participation is operational/site-specific.  A station that has
            # intentionally disabled one role stays disabled after template sync.
            merged["enabled"] = bool(local.get("enabled", merged.get("enabled", True)))
            used_current_ids.add(clean(local.get("id")))
        else:
            merged["path"] = ""
            merged["path_mode"] = "absolute"
            family = clean(merged.get("family_key"))
            suffix = clean(merged.get("family_suffix"))
            candidates = family_candidates(store, family, suffix=suffix) if family else ()
            if candidates:
                encoded, path_mode = encode_source_path(store, candidates[0])
                merged["path"] = encoded
                merged["path_mode"] = path_mode

        merged_sources.append(merged)

    # Preserve site-only extra sources.  They remain available and keep their
    # local settings, but do not silently become part of a profile rule unless
    # the reviewer explicitly maps them.
    for local in current.get("sources", []):
        local_id = clean(local.get("id"))
        if local_id and local_id not in used_current_ids:
            merged_sources.append(dict(local))

    merged_rules: list[dict] = []
    for rule in template.get("comparisons", []):
        bindings = {}
        for template_source_id, column_id in dict(rule.get("bindings") or {}).items():
            target_id = id_map.get(clean(template_source_id))
            if target_id and clean(column_id):
                bindings[target_id] = clean(column_id)
        merged_rules.append({
            "id": clean(rule.get("id")) or _rule_id(),
            "title": clean(rule.get("title")) or "Comparison",
            "comparison_mode": clean(rule.get("comparison_mode")) or COMPARISON_MODE_DEFAULT,
            "bindings": bindings,
        })

    result = normalize_config({
        "version": CONFIG_VERSION,
        "default_comparison_mode": _site_comparison_mode(template.get("default_comparison_mode")),
        "sources": merged_sources,
        "comparisons": merged_rules,
        "updated_at": clean(current.get("updated_at")),
    })
    metadata = {
        "profile_name": clean(profile.get("name")),
        "profile_modified_at": clean(profile.get("modified_at")),
        "applied_at": datetime.now().isoformat(timespec="seconds"),
    }
    return result, metadata


def source_column_catalog(store, source: dict) -> tuple[PhysicalColumn, ...]:
    path = source_path(store, source)
    if not path or not path.exists():
        return ()
    return inspect_table(
        path,
        sheet_name=clean(source.get("sheet_name")),
        header_row=max(1, int(source.get("header_row") or 1)),
    ).columns


def _path_is_inside_project_data(store, path: Path | None) -> bool:
    """Return True when *path* is an internal Project Data artifact.

    Configurable comparison sources are live links.  Historical releases copied
    repository files under ``project/source_files``; v0.8.182 must never prefer
    those archived snapshots when an original/live file can be resolved.
    """
    if path is None:
        return False
    try:
        Path(path).resolve().relative_to(Path(store.folder).resolve())
        return True
    except (OSError, ValueError):
        return False


def _find_site_file_by_name(store, file_name: str) -> Path | None:
    wanted = Path(str(file_name or "")).name.casefold()
    if not wanted:
        return None
    matches = [path for path in scan_site_tabular_files(store) if path.name.casefold() == wanted]
    if not matches:
        return None
    def rank(path: Path):
        try:
            return (int(path.stat().st_mtime_ns), str(path).casefold())
        except OSError:
            return (0, str(path).casefold())
    return max(matches, key=rank)


def _latest_import_original_name(store, source_type: str) -> str:
    try:
        row = store.db.execute(
            "SELECT original_name FROM imports WHERE source_type=? ORDER BY id DESC LIMIT 1",
            (clean(source_type),),
        ).fetchone()
        return clean(row[0]) if row else ""
    except Exception:
        return ""


def _legacy_path(store, source_type: str) -> Path | None:
    """Resolve a legacy role to its live/original file without making a copy.

    Resolution order deliberately favors reviewer-facing files: explicit manual
    original path, last remembered live path, the original basename recorded by
    the legacy importer, then (only as a final compatibility fallback) the old
    internal snapshot.  New configurable saves encode the resolved live path.
    """
    try:
        manual = (store.manual_source_overrides().get(source_type) or {}) if hasattr(store, "manual_source_overrides") else {}
        original = str(manual.get("original_path") or "").strip()
        if original:
            candidate = Path(original)
            if candidate.exists() and candidate.is_file():
                return candidate

        live_meta = dict((store.config or {}).get("live_source_metadata", {}) or {})
        meta = dict(live_meta.get(source_type) or {})
        live_text = str(meta.get("path") or "").strip()
        if live_text:
            candidate = Path(live_text)
            if candidate.exists() and candidate.is_file():
                return candidate

        original_name = _latest_import_original_name(store, source_type)
        if original_name:
            candidate = _find_site_file_by_name(store, original_name)
            if candidate is not None:
                return candidate

        path = store.source_path(source_type) if hasattr(store, "source_path") else None
        return Path(path) if path else None
    except Exception:
        return None


def _guess_key_column(columns: tuple[PhysicalColumn, ...]) -> str:
    preferred = (
        "equipment name", "devicename", "device_name", "de_name", "equipment",
        "rmu_name", "rmu_no", "rmu", "name", "实例名称",
    )
    by = {column.header.strip().casefold(): column.id for column in columns}
    for token in preferred:
        if token.casefold() in by:
            return by[token.casefold()]
    return columns[0].id if columns else ""


def bootstrap_legacy_sources(store) -> dict:
    """One-time compatibility bootstrap for existing v0.8.177 projects.

    The resulting configuration is no longer fixed: it is simply an editable
    starting point made from whatever legacy files happen to be active.
    """
    labels = (
        ("se_list", "SE"), ("zenon_db", "ZENON DB"), ("zenon_sld", "ZENON SLD"),
        ("adms_db", "ADMS DB"), ("adms_sld", "ADMS SLD"),
    )
    config = empty_config()
    for source_type, label in labels:
        path = _legacy_path(store, source_type)
        if not path or not path.exists():
            continue
        sheet_name = ""
        try:
            if path.suffix.lower() in {".xlsx", ".xlsm"} and hasattr(store, "source_sheet_name"):
                sheet_name = clean(store.source_sheet_name(source_type))
            structure = inspect_table(path, sheet_name=sheet_name)
        except Exception:
            continue
        encoded, mode = encode_source_path(store, path)
        config["sources"].append({
            "id": _source_id(),
            "title": label or path.stem,
            "path": encoded,
            "path_mode": mode,
            "sheet_name": structure.sheet_name,
            "header_row": structure.header_row,
            "key_column": _guess_key_column(structure.columns),
            "hidden_columns": [],
            "enabled": True,
            "selection_mode": "pinned",
            "family_key": file_family_key(path),
            "family_suffix": path.suffix.lower(),
        })
    # This is an editable migration proposal only.  Merely opening the
    # configuration dialog must never mutate an existing v0.8.177 site.  The
    # proposal is persisted only when the reviewer presses Save.
    return config


def _file_context(path: Path, title: str) -> dict:
    """Compact source provenance used by the existing lifecycle audit.

    Configurable sources do not have legacy SE/ZENON/ADMS role ids, therefore
    we keep the same source-change lifecycle by attaching provenance under the
    stable configurable source id.
    """
    path = Path(path)
    try:
        stat = path.stat()
    except OSError:
        return {"name": path.name, "title": clean(title) or path.stem}
    digest = hashlib.sha256()
    try:
        with path.open("rb") as handle:
            for chunk in iter(lambda: handle.read(1024 * 1024), b""):
                digest.update(chunk)
        sha256 = digest.hexdigest()
    except OSError:
        sha256 = ""
    return {
        "name": path.name,
        "title": clean(title) or path.stem,
        "path": str(path),
        "size": int(stat.st_size),
        "mtime_ns": int(stat.st_mtime_ns),
        "sha256": sha256,
    }


def config_status(store, config: dict | None = None) -> list[dict]:
    config = normalize_config(config or get_config(store))
    rule_bindings = {source["id"]: 0 for source in config["sources"]}
    for rule in config["comparisons"]:
        for source_id in rule.get("bindings", {}):
            if source_id in rule_bindings:
                rule_bindings[source_id] += 1
    result = []
    for source in config["sources"]:
        path = source_path(store, source)
        row = {"source": source, "path": path, "status": "READY", "columns": (), "error": ""}
        if not bool(source.get("enabled", True)):
            row["status"] = "DISABLED"
        elif not path or not path.exists():
            row["status"] = "MISSING"
        else:
            try:
                columns = source_column_catalog(store, source)
                row["columns"] = columns
                ids = {column.id for column in columns}
                if not clean(source.get("key_column")) or source.get("key_column") not in ids:
                    row["status"] = "KEY REQUIRED"
            except Exception as exc:
                row["status"] = "ERROR"
                row["error"] = f"{type(exc).__name__}: {exc}"
        row["compare_count"] = rule_bindings.get(source["id"], 0)
        result.append(row)
    return result


def _generic_normalize(value: object) -> str:
    """Generic equality normalizer: trim only, plus stable Excel-number text.

    No business-specific FEEDER/SMART/TYPE assumptions are applied because the
    source is intentionally fully configurable.  What the reviewer maps is what
    the App compares.
    """
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return format(value, ".15g")
    return clean(value)


def _index_rows(table: LoadedTable, key_column: str) -> tuple[dict[str, dict], dict[str, int], dict[str, str]]:
    grouped: dict[str, list[dict]] = {}
    display: dict[str, str] = {}
    for row in table.rows:
        raw_key = _generic_normalize(row.get(key_column))
        if not raw_key:
            continue
        normalized_key = raw_key.casefold()
        grouped.setdefault(normalized_key, []).append(row)
        display.setdefault(normalized_key, raw_key)
    return (
        {key: rows[0] for key, rows in grouped.items()},
        {key: len(rows) for key, rows in grouped.items()},
        display,
    )


def review_column_key(source_id: str, column_id: str) -> str:
    safe_source = re.sub(r"[^a-zA-Z0-9_]+", "_", clean(source_id))
    safe_column = re.sub(r"[^a-zA-Z0-9_]+", "_", clean(column_id))
    return f"generic__{safe_source}__{safe_column}"


def analysis_column_key(rule_id: str) -> str:
    safe = re.sub(r"[^a-zA-Z0-9_]+", "_", clean(rule_id))
    return f"analysis__{safe}"


def analysis_detail_key(rule_id: str) -> str:
    return analysis_column_key(rule_id) + "__detail"


def _analysis_tooltip(title: str, result, mode: str, values_by_source: dict[str, object]) -> str:
    effective = _site_comparison_mode(mode)
    if effective == COMPARISON_MODE_IGNORE_BLANK:
        rule_text = (
            "Mode: Ignore blank values. Blank values are excluded; one remaining non-blank value is TRUE; "
            "different remaining values are FALSE; if every bound value is blank the result is N/A."
        )
    else:
        rule_text = (
            "Mode: Strict equality. Every bound source participates and blank is a real comparison value. "
            "All values equal = TRUE (including all blank); any blank/non-blank mix or other difference = FALSE."
        )
    lines = [f"{title} consistency check", rule_text, ""]
    if not values_by_source:
        lines += ["No bound sources", "Result: N/A"]
        return "\n".join(lines)
    for source, raw_value in values_by_source.items():
        raw = clean(raw_value)
        normalized = _generic_normalize(raw_value)
        lines.append(f"{source}: {raw or normalized or '<blank>'}")
    lines += ["", f"Result: {result.display or 'N/A'}"]
    return "\n".join(lines)


def build_configurable_review(store) -> tuple[list[dict], dict]:
    config = get_config(store, bootstrap=False)
    if not config["sources"]:
        return [], {"coverage": {}, "source_count": 0, "config": config}

    loaded: dict[str, LoadedTable] = {}
    indexes: dict[str, dict[str, dict]] = {}
    duplicate_counts: dict[str, dict[str, int]] = {}
    display_keys: dict[str, dict[str, str]] = {}
    source_meta: dict[str, dict] = {}

    for source in config["sources"]:
        if not bool(source.get("enabled", True)):
            continue
        source_id = source["id"]
        path = source_path(store, source)
        if not path or not path.exists():
            continue
        table = load_table(
            path,
            sheet_name=clean(source.get("sheet_name")),
            header_row=max(1, int(source.get("header_row") or 1)),
        )
        key_column = clean(source.get("key_column"))
        available_ids = {column.id for column in table.structure.columns}
        if not key_column or key_column not in available_ids:
            continue
        loaded[source_id] = table
        index, duplicates, display = _index_rows(table, key_column)
        indexes[source_id] = index
        duplicate_counts[source_id] = duplicates
        display_keys[source_id] = display
        source_title = clean(source.get("title")) or path.stem
        source_meta[source_id] = {
            "title": source_title,
            "path": str(path),
            "columns": {column.id: column for column in table.structure.columns},
            "hidden": set(source.get("hidden_columns") or []),
            "context": _file_context(path, source_title),
        }

    all_keys: set[str] = set()
    for index in indexes.values():
        all_keys.update(index)

    output: list[dict] = []
    coverage = {}
    ordered_sources = [source for source in config["sources"] if bool(source.get("enabled", True)) and source["id"] in loaded]
    ordered_rules = list(config["comparisons"])

    for normalized_key in sorted(all_keys):
        present_sources = [source for source in ordered_sources if normalized_key in indexes.get(source["id"], {})]
        # Prefer the first configured source's original spelling for the Row Locator.
        display_key = next(
            (display_keys[source["id"]].get(normalized_key) for source in ordered_sources if display_keys[source["id"]].get(normalized_key)),
            normalized_key,
        )
        present_count = len(present_sources)
        coverage[f"{present_count}/{max(1, len(ordered_sources))}"] = coverage.get(f"{present_count}/{max(1, len(ordered_sources))}", 0) + 1
        missing_titles = [
            (source_meta[source["id"]]["title"])
            for source in ordered_sources if normalized_key not in indexes.get(source["id"], {})
        ]

        row: dict = {
            "no": len(output) + 1,
            "rmu": display_key,
            "review_key": display_key,
            "equipment_device_type": "EQUIPMENT",
            "equipment_source_count": f"{present_count}/{len(ordered_sources)}" if ordered_sources else "0/0",
            "equipment_missing_sources": ", ".join(missing_titles),
            "equipment_source_presence_detail": "Present: " + ", ".join(source_meta[s["id"]]["title"] for s in present_sources)
                + ("\nMissing: " + ", ".join(missing_titles) if missing_titles else ""),
            "comments": "",
            "remarks": "",
            "status": "MATCHED",
            "analysis_field_order": [],
            "analysis_field_labels": {},
            "analysis_field_keys": {},
            "resolution_candidates": {},
            # Keep the existing Needs Action lifecycle/source-change audit alive
            # even though source roles are now arbitrary and site-local.
            "_source_context": {
                source["id"]: dict(source_meta[source["id"]].get("context") or {})
                for source in ordered_sources
            },
            "_source_audit_names": {
                source["id"]: source_meta[source["id"]]["title"]
                for source in ordered_sources
            },
            "_source_field_labels": {
                source["id"]: {
                    column.id: column.label for column in loaded[source["id"]].structure.columns
                }
                for source in ordered_sources
            },
            "_comparison_config_version": CONFIG_VERSION,
        }

        duplicate_messages = []
        for source in ordered_sources:
            source_id = source["id"]
            source_row = indexes.get(source_id, {}).get(normalized_key, {})
            if duplicate_counts.get(source_id, {}).get(normalized_key, 0) > 1:
                duplicate_messages.append(
                    f"{source_meta[source_id]['title']}: {duplicate_counts[source_id][normalized_key]} rows share this key; first row is used"
                )
            for column in loaded[source_id].structure.columns:
                row[review_column_key(source_id, column.id)] = clean(source_row.get(column.id)) if source_row else ""

        mismatches = []
        for rule in ordered_rules:
            rule_id = rule["id"]
            title = clean(rule.get("title")) or "Comparison"
            values = {}
            for source in ordered_sources:
                source_id = source["id"]
                column_id = clean((rule.get("bindings") or {}).get(source_id))
                if not column_id:
                    # No binding means this source does not participate in this rule.
                    continue
                source_row = indexes.get(source_id, {}).get(normalized_key, {})
                # v0.8.191 strict comparison: a bound source always participates.
                # Missing source rows and blank physical fields therefore contribute
                # an explicit blank value instead of being silently ignored.
                values[source_meta[source_id]["title"]] = source_row.get(column_id, "") if source_row else ""
            effective_mode = _rule_comparison_mode(rule, config.get("default_comparison_mode"))
            if effective_mode == COMPARISON_MODE_IGNORE_BLANK:
                result = compare_consistency(values, _generic_normalize)
            else:
                result = compare_strict_consistency(values, _generic_normalize)
            analysis_key = analysis_column_key(rule_id)
            detail_key = analysis_detail_key(rule_id)
            row[analysis_key] = result.display
            row[detail_key] = _analysis_tooltip(title, result, effective_mode, values)
            row["analysis_field_order"].append(rule_id)
            row["analysis_field_labels"][rule_id] = title
            row["analysis_field_keys"][rule_id] = analysis_key
            row["resolution_candidates"][rule_id] = [
                {
                    "source": source,
                    "value": clean((result.raw_by_source or {}).get(source) or normalized),
                    "normalized": clean(normalized),
                }
                for source, normalized in (result.normalized_by_source or {}).items()
            ]
            if result.value is False:
                mismatches.append(title)

        remarks = []
        if missing_titles:
            remarks.append("Missing source row: " + ", ".join(missing_titles))
        if mismatches:
            remarks.append("Analysis mismatch: " + " / ".join(mismatches))
            row["status"] = "WARNING"
        elif ordered_rules:
            remarks.append("Analysis: all available configured checks passed")
        remarks.extend(duplicate_messages)
        row["remarks"] = "; ".join(remarks)
        output.append(row)

    return output, {
        "coverage": coverage,
        "source_count": len(ordered_sources),
        "config": config,
        "source_meta": source_meta,
    }


def review_groups(store) -> tuple:
    """Dynamic Equipment Data Review groups for the active site configuration."""
    config = get_config(store, bootstrap=False)
    groups: list[tuple[str, str, tuple]] = []
    groups.append(("Index", "#FFFFFF", (
        ("no", "No.", 55), ("rmu", "Index / Key", 165),
    )))
    groups.append(("Source Coverage", "#EEF2F6", (
        ("equipment_source_count", "Sources", 90),
        ("equipment_missing_sources", "Missing Sources", 260),
    )))
    analysis_columns = tuple(
        (analysis_column_key(rule["id"]), clean(rule.get("title")) or "Comparison", 100)
        for rule in config["comparisons"]
    )
    if analysis_columns:
        groups.append(("Analysis", "#EE94A0", analysis_columns))
    groups.append(("Remarks", "#D5D8DE", (("remarks", "Remarks", 300),)))
    groups.append(("Resolution", "#D8E1EA", (("comments", "Resolution / Comments", 360),)))

    palette = ("#83D9D2", "#E3F2D9", "#C1E3AC", "#91ACDF", "#B6C7EA", "#EADCF5", "#F5E3C6", "#D8ECEA")
    visible_sources = [source for source in config["sources"] if bool(source.get("enabled", True))]
    for source_index, source in enumerate(visible_sources):
        path = source_path(store, source)
        title = clean(source.get("title")) or (path.stem if path else "Source")
        hidden = set(source.get("hidden_columns") or [])
        columns = []
        try:
            catalog = source_column_catalog(store, source)
        except Exception:
            catalog = ()
        for column in catalog:
            if column.id in hidden:
                continue
            columns.append((review_column_key(source["id"], column.id), column.label, 150))
        groups.append((title, palette[source_index % len(palette)], tuple(columns)))
    return tuple(groups)


def configured_analysis_fields(data: dict) -> tuple[tuple[str, str, str], ...]:
    """Return (internal_id, display_label, row_value_key) for one review row."""
    order = list(data.get("analysis_field_order") or [])
    labels = dict(data.get("analysis_field_labels") or {})
    keys = dict(data.get("analysis_field_keys") or {})
    return tuple((field_id, clean(labels.get(field_id)) or field_id, clean(keys.get(field_id)) or analysis_column_key(field_id)) for field_id in order)
