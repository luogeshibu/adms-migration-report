"""Configurable source mapping, joins and derived-table generation.

This module is deliberately isolated from the proven RMU/Signal validation
engines.  It consumes the same mapped source tables but does not change their
business rules.  Users can define extra source fields, join tables, calculate
new fields, preview results and export a new CSV/XLSX table.
"""
from __future__ import annotations

import ast
import csv
import json
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Callable, Mapping

from openpyxl import Workbook

from ..config.sources import schema_for
from ..infrastructure.parsers import read_csv_raw, read_excel_raw, read_mapped_rows, read_standard_raw, resolve_source_excel_sheet_name
from ..parsers import clean
from .schema_service import get_source_overrides


SOURCE_ALIASES: dict[str, str] = {
    "se_list": "se",
    "zenon_db": "zenon_db",
    "zenon_sld": "zenon_sld",
    "adms_db": "adms_db",
    "adms_sld": "adms_sld",
    "ioa": "ioa",
    "standard_reference": "standard",
}

SOURCE_LABELS: dict[str, str] = {
    "se_list": "SE",
    "zenon_db": "ZENON DB",
    "zenon_sld": "ZENON SLD",
    "adms_db": "ADMS DB",
    "adms_sld": "ADMS SLD",
    "ioa": "ZENON-ADMS IOA",
    "standard_reference": "IOA STANDARD",
}


def source_alias(source_type: str) -> str:
    return SOURCE_ALIASES.get(str(source_type), re.sub(r"\W+", "_", str(source_type)).strip("_").lower())


def source_label(source_type: str) -> str:
    schema = schema_for(source_type)
    return SOURCE_LABELS.get(source_type, schema.label if schema else str(source_type))


def available_source_types() -> tuple[str, ...]:
    return tuple(SOURCE_ALIASES)


def slug_key(value: str, fallback: str = "field") -> str:
    text = re.sub(r"[^0-9A-Za-z_]+", "_", clean(value)).strip("_").lower()
    if not text:
        text = fallback
    if text[0].isdigit():
        text = "f_" + text
    return text


def _raw_rows(source_type: str, path: Path, overrides: Mapping[str, str] | None = None, *, sheet_name: str | None = None):
    schema = schema_for(source_type)
    path = Path(path)
    if source_type == "standard_reference":
        return read_standard_raw(
            path,
            sheet_name=(schema.sheet_name if schema and schema.sheet_name else "STANDARD"),
            scan_rows=(schema.header_scan_rows if schema else 25),
            overrides=overrides,
        )
    if path.suffix.lower() == ".csv":
        return read_csv_raw(path)
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        return read_excel_raw(path, sheet_name=sheet_name)
    raise ValueError(f"Unsupported derived-table source format: {path.name}")


def load_source_rows(store, source_type: str, path: Path) -> list[dict]:
    """Return canonical built-in fields plus persistent custom source fields."""
    overrides = get_source_overrides(store, source_type)
    sheet_name = store.source_sheet_name(source_type) if store and Path(path).suffix.lower() in {".xlsx", ".xlsm"} and hasattr(store, "source_sheet_name") else ""
    effective_sheet = resolve_source_excel_sheet_name(source_type, Path(path), sheet_name or None) if Path(path).suffix.lower() in {".xlsx", ".xlsm"} else ""
    mapped = read_mapped_rows(source_type, Path(path), overrides, strict=False, sheet_name=effective_sheet or None)
    _headers, raw = _raw_rows(source_type, Path(path), overrides, sheet_name=effective_sheet or None)
    custom = list(store.custom_source_fields(source_type) if store else [])
    rows: list[dict] = []
    count = max(len(mapped.rows), len(raw))
    for index in range(count):
        canonical = dict(mapped.rows[index]) if index < len(mapped.rows) else {}
        original = dict(raw[index]) if index < len(raw) else {}
        for field in custom:
            key = clean(field.get("field_key"))
            actual = clean(field.get("actual_column"))
            if key:
                canonical[key] = original.get(actual) if actual else None
        rows.append(canonical)
    return rows


def source_field_catalog(store, source_type: str) -> list[dict]:
    schema = schema_for(source_type)
    result: list[dict] = []
    if schema:
        for spec in schema.fields:
            result.append({"key": spec.key, "label": spec.label, "custom": False})
    if store:
        for item in store.custom_source_fields(source_type):
            result.append({
                "key": clean(item.get("field_key")),
                "label": clean(item.get("display_name")) or clean(item.get("field_key")),
                "custom": True,
            })
    return [item for item in result if item["key"]]


class _RecordProxy:
    def __init__(self, row: Mapping[str, object] | None):
        self._row = dict(row or {})

    def field(self, name: str):
        return self._row.get(name)


class SafeExpressionEvaluator:
    """Tiny safe expression language used by derived output fields.

    Supported examples::

        adms_db.rmu
        COALESCE(adms_db.smart, adms_sld.smart, se.smart)
        IF(NORMALIZE(zenon_sld.feeder) == NORMALIZE(adms_db.gss_fid), "Closed", "Needs Action")
        CONCAT(se.station, "-", adms_db.rmu)

    No Python eval/exec is used.
    """

    def __init__(self, context: Mapping[str, Mapping[str, object] | None]):
        self.context = {str(k): _RecordProxy(v) for k, v in context.items()}

    @staticmethod
    def normalize(value):
        text = clean(value).upper()
        return re.sub(r"[^0-9A-Z]+", "", text)

    @staticmethod
    def coalesce(*values):
        for value in values:
            if value is not None and clean(value) != "":
                return value
        return ""

    @staticmethod
    def concat(*values):
        return "".join("" if value is None else str(value) for value in values)

    @staticmethod
    def upper(value):
        return clean(value).upper()

    @staticmethod
    def lower(value):
        return clean(value).lower()

    @staticmethod
    def contains(value, needle):
        return clean(needle).casefold() in clean(value).casefold()

    def evaluate(self, expression: str):
        expression = clean(expression)
        if not expression:
            return ""
        try:
            tree = ast.parse(expression, mode="eval")
        except SyntaxError as exc:
            raise ValueError(f"Invalid expression: {exc.msg}") from exc
        return self._eval(tree.body)

    def _eval(self, node):
        if isinstance(node, ast.Constant):
            return node.value
        if isinstance(node, ast.Name):
            name = node.id
            if name in self.context:
                return self.context[name]
            if name in {"True", "False", "None"}:
                return {"True": True, "False": False, "None": None}[name]
            raise ValueError(f"Unknown name '{name}'")
        if isinstance(node, ast.Attribute):
            base = self._eval(node.value)
            if isinstance(base, _RecordProxy):
                return base.field(node.attr)
            raise ValueError("Only source.field references are allowed")
        if isinstance(node, ast.Call):
            if not isinstance(node.func, ast.Name):
                raise ValueError("Only named functions are allowed")
            name = node.func.id.upper()
            args = [self._eval(arg) for arg in node.args]
            funcs = {
                "COALESCE": self.coalesce,
                "IF": lambda condition, yes, no="": yes if bool(condition) else no,
                "NORMALIZE": self.normalize,
                "CONCAT": self.concat,
                "UPPER": self.upper,
                "LOWER": self.lower,
                "CONTAINS": self.contains,
                "STR": lambda value: "" if value is None else str(value),
                "INT": lambda value: int(float(value)) if clean(value) else 0,
                "FLOAT": lambda value: float(value) if clean(value) else 0.0,
            }
            if name not in funcs:
                raise ValueError(f"Unsupported function '{name}'")
            return funcs[name](*args)
        if isinstance(node, ast.Compare):
            left = self._eval(node.left)
            for operator, comparator in zip(node.ops, node.comparators):
                right = self._eval(comparator)
                if isinstance(operator, ast.Eq): ok = left == right
                elif isinstance(operator, ast.NotEq): ok = left != right
                elif isinstance(operator, ast.Lt): ok = left < right
                elif isinstance(operator, ast.LtE): ok = left <= right
                elif isinstance(operator, ast.Gt): ok = left > right
                elif isinstance(operator, ast.GtE): ok = left >= right
                elif isinstance(operator, ast.In): ok = left in right
                elif isinstance(operator, ast.NotIn): ok = left not in right
                else: raise ValueError("Unsupported comparison")
                if not ok:
                    return False
                left = right
            return True
        if isinstance(node, ast.BoolOp):
            values = [bool(self._eval(value)) for value in node.values]
            if isinstance(node.op, ast.And): return all(values)
            if isinstance(node.op, ast.Or): return any(values)
            raise ValueError("Unsupported boolean operator")
        if isinstance(node, ast.UnaryOp):
            value = self._eval(node.operand)
            if isinstance(node.op, ast.Not): return not bool(value)
            if isinstance(node.op, ast.USub): return -value
            raise ValueError("Unsupported unary operator")
        if isinstance(node, ast.BinOp):
            left, right = self._eval(node.left), self._eval(node.right)
            if isinstance(node.op, ast.Add): return left + right
            if isinstance(node.op, ast.Sub): return left - right
            if isinstance(node.op, ast.Mult): return left * right
            if isinstance(node.op, ast.Div): return left / right
            raise ValueError("Unsupported arithmetic operator")
        raise ValueError(f"Unsupported expression element: {type(node).__name__}")


def evaluate_expression(expression: str, context: Mapping[str, Mapping[str, object] | None]):
    return SafeExpressionEvaluator(context).evaluate(expression)


def _join_key(value) -> str:
    return clean(value).strip().casefold()


def _context_with_alias(source_type: str, row: dict | None) -> dict:
    return {source_alias(source_type): row}


@dataclass(frozen=True)
class DerivedBuildResult:
    columns: tuple[str, ...]
    rows: tuple[dict, ...]
    warnings: tuple[str, ...] = ()


def build_derived_table(
    config: Mapping[str, object],
    store,
    source_paths: Mapping[str, Path | str | None],
) -> DerivedBuildResult:
    base_source = clean(config.get("base_source_type"))
    if not base_source:
        raise ValueError("Select a Base Source first.")
    base_path = source_paths.get(base_source)
    if not base_path or not Path(base_path).exists():
        raise ValueError(f"Base Source is unavailable: {source_label(base_source)}")

    cache: dict[str, list[dict]] = {}

    def rows_for(source_type: str) -> list[dict]:
        if source_type not in cache:
            path = source_paths.get(source_type)
            if not path or not Path(path).exists():
                raise ValueError(f"Source is unavailable: {source_label(source_type)}")
            cache[source_type] = load_source_rows(store, source_type, Path(path))
        return cache[source_type]

    contexts = [_context_with_alias(base_source, row) for row in rows_for(base_source)]
    warnings: list[str] = []
    joined_sources = {base_source}

    for join in list(config.get("joins") or []):
        source_type = clean(join.get("source_type"))
        if not source_type or source_type in joined_sources:
            continue
        join_type = clean(join.get("join_type") or "LEFT").upper()
        if join_type not in {"LEFT", "INNER", "FULL"}:
            raise ValueError(f"Unsupported join type: {join_type}")
        left_expression = clean(join.get("left_expression"))
        right_field = clean(join.get("right_field"))
        if not left_expression or not right_field:
            raise ValueError(f"Join for {source_label(source_type)} requires both join fields.")

        target_rows = rows_for(source_type)
        index: dict[str, list[tuple[int, dict]]] = {}
        for pos, row in enumerate(target_rows):
            index.setdefault(_join_key(row.get(right_field)), []).append((pos, row))
        matched_positions: set[int] = set()
        new_contexts: list[dict] = []
        alias = source_alias(source_type)
        for context in contexts:
            try:
                left_value = evaluate_expression(left_expression, context)
            except Exception as exc:
                raise ValueError(f"Join expression '{left_expression}' failed: {exc}") from exc
            matches = index.get(_join_key(left_value), []) if clean(left_value) else []
            if matches:
                for pos, row in matches:
                    matched_positions.add(pos)
                    merged = dict(context)
                    merged[alias] = row
                    new_contexts.append(merged)
            elif join_type in {"LEFT", "FULL"}:
                merged = dict(context)
                merged[alias] = None
                new_contexts.append(merged)
        if join_type == "FULL":
            known_aliases = [source_alias(item) for item in joined_sources]
            for pos, row in enumerate(target_rows):
                if pos in matched_positions:
                    continue
                context = {known_alias: None for known_alias in known_aliases}
                context[alias] = row
                new_contexts.append(context)
        contexts = new_contexts
        joined_sources.add(source_type)

    fields = list(config.get("fields") or [])
    if not fields:
        raise ValueError("Add at least one Output Field.")
    columns: list[str] = []
    for index, field in enumerate(fields, start=1):
        name = clean(field.get("name")) or f"Field {index}"
        columns.append(name)

    output_rows: list[dict] = []
    for row_no, context in enumerate(contexts, start=1):
        output: dict[str, object] = {}
        for index, field in enumerate(fields):
            name = columns[index]
            expression = clean(field.get("expression"))
            try:
                value = evaluate_expression(expression, context)
            except Exception as exc:
                value = ""
                warning = f"Row {row_no}, {name}: {exc}"
                if len(warnings) < 100:
                    warnings.append(warning)
            output[name] = value
        output_rows.append(output)
    return DerivedBuildResult(tuple(columns), tuple(output_rows), tuple(warnings))


def export_derived_csv(result: DerivedBuildResult, path: Path) -> Path:
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    with target.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(result.columns), extrasaction="ignore")
        writer.writeheader()
        writer.writerows(result.rows)
    return target


def export_derived_xlsx(result: DerivedBuildResult, path: Path) -> Path:
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "DATA"
    for col, name in enumerate(result.columns, start=1):
        ws.cell(1, col).value = name
    for r, row in enumerate(result.rows, start=2):
        for c, name in enumerate(result.columns, start=1):
            ws.cell(r, c).value = row.get(name)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for index, name in enumerate(result.columns, start=1):
        width = max(12, min(45, len(str(name)) + 4))
        ws.column_dimensions[ws.cell(1, index).column_letter].width = width
    wb.save(target)
    wb.close()
    return target


def default_derived_config(name: str = "RMU_FINAL") -> dict:
    return {
        "table_name": clean(name) or "RMU_FINAL",
        "description": "",
        "base_source_type": "adms_db",
        "joins": [],
        "fields": [
            {"name": "RMU", "key": "rmu", "expression": "adms_db.rmu"},
        ],
    }


def config_json(config: Mapping[str, object]) -> str:
    return json.dumps(dict(config), ensure_ascii=False, sort_keys=True, indent=2)
